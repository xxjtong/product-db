"""Product API routes — v2 schema with dictionary + mapping tables."""
from __future__ import annotations
import json
import logging
import os
import re
import subprocess
import tempfile
import uuid
from typing import Optional
from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, File
from fastapi.responses import StreamingResponse, HTMLResponse, Response
from sqlalchemy.orm import Session, joinedload, selectinload
from sqlalchemy import select, text, case
from sqlalchemy import or_
from app.database import get_db
from app.utils.helpers import get_or_404, apply_partial_update, format_description_with_specs
from app.models.product import Product
from app.models.category import Category, CategorySpecDefinition
from app.models.dependency import ProductDependency
from app.models.supplier import Supplier
from app.models.dictionary import Manufacturer
from app.models.mapping import (ProductCommMethod, ProductCommProtocol, ProductPowerSupply,
                                ProductHardwareInterface, ProductSensorCapability, ProductImage)
from app.services.spec_service import validate_specs, compare_products
from app.services.storage import save_upload, delete_file, upload_from_url
from app.services.spec_generator import build_spec_html
from app.services.ai_extract import call_ai_extract, regex_extract_from_text
from app.services.product_helpers import (
    build_pinyin, get_name_maps, product_eager_loads, build_product_detail,
    write_mappings, rewrite_mappings,
)
from app.auth import get_current_user, filter_by_ownership, check_ownership
from app.config import settings
from app.models.user import User
from app.utils.escape import escape_like
from app.utils.security import validate_url
from app.schemas.product import (
    ProductCreate, ProductUpdate, AIFetchRequest,
    DownloadImageRequest,
    ProductDependencyCreate, ProductDependencyUpdate,
)
import openpyxl
import httpx
from bs4 import BeautifulSoup
from urllib.parse import urlparse
from io import BytesIO
from datetime import datetime

router = APIRouter()

# --- Static routes (must be before {product_id}) ---

@router.get("/products")
def list_products(
    category_id: Optional[int] = None,
    search: str = "",
    parent_id: Optional[int] = Query(None),
    comm_method: Optional[int] = Query(None, description="Filter by comm method dict ID"),
    comm_protocol: Optional[int] = Query(None, description="Filter by protocol dict ID"),
    power_supply: Optional[int] = Query(None, description="Filter by power supply dict ID"),
    manufacturer_id: Optional[int] = None,
    status: str = "active",
    page: int = 1,
    per_page: int = 20,
    db: Session = Depends(get_db),
    user=Depends(get_current_user),
):
    q = db.query(Product).options(*product_eager_loads())
    q = filter_by_ownership(q, Product, user)

    if category_id:
        from app.services.product_category_helper import get_products_in_categories, get_category_descendants
        cat_ids = get_category_descendants(db, category_id)
        m2m_ids = get_products_in_categories(db, cat_ids)
        q = q.filter(
            or_(
                Product.id.in_(m2m_ids),
                Product.category_id.in_(cat_ids)  # fallback: old single-category column
            )
        )

    if search:
        search_lower = search.lower()
        q = q.filter(
            or_(
                Product.name.ilike(f"%{escape_like(search)}%"),
                Product.model.ilike(f"%{escape_like(search)}%"),
                Product.sku.ilike(f"%{escape_like(search)}%"),
                Product.pinyin_search.ilike(f"%{escape_like(search_lower)}%"),
                Product.description.ilike(f"%{escape_like(search)}%"),
                # Also search by category name via product_categories junction
                Product.id.in_(
                    select(text('pc.product_id')).select_from(text('product_categories pc JOIN device_categories c ON pc.category_id = c.id')).where(text('c.name LIKE :s')).params(s=f'%{escape_like(search)}%')
                ),
            )
        )

    if parent_id is not None:
        if parent_id == 0:
            q = q.filter(Product.parent_id.is_(None))
        else:
            q = q.filter(Product.parent_id == parent_id)

    if status:
        q = q.filter(Product.status == status)

    if manufacturer_id:
        q = q.filter(Product.manufacturer_id == manufacturer_id)

    # Mapping table filters
    if comm_method:
        q = q.filter(Product.comm_methods.any(ProductCommMethod.method_id == comm_method))
    if comm_protocol:
        q = q.filter(Product.comm_protocols.any(ProductCommProtocol.protocol_id == comm_protocol))
    if power_supply:
        q = q.filter(Product.power_supplies.any(ProductPowerSupply.power_id == power_supply))

    total = q.count()
    # Order: priority manufacturers (sort_order < 100) → products with images → others
    _mfg_sort_map = {m.id: m.sort_order for m in db.query(Manufacturer).filter(Manufacturer.sort_order < 100).all()}
    _mfg_priority_ids = list(_mfg_sort_map.keys())
    _sort_key = case(
        (Product.manufacturer_id.in_(_mfg_priority_ids), 0),
        (Product.image_url != None, 1),
        else_=2,
    )
    # Within priority group, sort by manufacturer sort_order then name
    _mfg_cases = [(Product.manufacturer_id == mid, so) for mid, so in _mfg_sort_map.items()]
    if _mfg_cases:
        _mfg_detail_key = case(*_mfg_cases, else_=999)
        products = q.order_by(_sort_key, _mfg_detail_key, Product.name).offset((page - 1) * per_page).limit(per_page).all()
    else:
        products = q.order_by(_sort_key, Product.name).offset((page - 1) * per_page).limit(per_page).all()

    cats, mfgs, sups = get_name_maps(db)

    product_list = [p.to_dict(cats, sups, mfgs) for p in products]
    # Enrich with multi-category names
    from app.models.category import Category as CatModel
    from app.services.product_category_helper import get_product_category_map
    cat_id_to_name = {c.id: c.name for c in db.query(CatModel).all()}
    pc_id_map = get_product_category_map(db)
    pc_map: dict[int, list[str]] = {}
    for pid, cids in pc_id_map.items():
        pc_map[pid] = [cat_id_to_name[cid] for cid in cids if cid in cat_id_to_name]
    for p in product_list:
        pid = p['id']
        extra = pc_map.get(pid, [])
        if extra: p['all_category_names'] = extra
    # Apply field visibility for non-admin users
    from app.services.field_visibility import filter_fields_for_user
    is_admin = getattr(user, 'role', '') == 'admin'
    for p in product_list:
        filter_fields_for_user(p, is_admin, db)
    return {
        "products": product_list,
        "total": total,
        "page": page,
        "per_page": per_page,
    }


@router.get("/products/compare")
def compare(product_ids: str, db: Session = Depends(get_db), user=Depends(get_current_user)):
    ids = [int(x) for x in product_ids.split(",") if x.strip().isdigit()]
    if len(ids) < 2:
        raise HTTPException(400, "At least 2 product IDs required")
    products = db.query(Product).options(*product_eager_loads()).filter(Product.id.in_(ids)).all()
    if len(products) < 2:
        raise HTTPException(404, "Products not found")

    matrix = compare_products(products)
    product_data = {
        p.id: {
            "id": p.id,
            "name": p.name,
            "model": p.model or "",
            "sku": p.sku or "",
            "manufacturer_name": p.manufacturer.name if p.manufacturer else "",
            "category_name": p.category.name if p.category else "",
        }
        for p in products
    }

    categories = set(p.category_id for p in products)
    spec_defs = db.query(CategorySpecDefinition).filter(
        CategorySpecDefinition.category_id.in_(categories)
    ).all()
    display_map = {sd.spec_key: sd.display_name for sd in spec_defs}

    return {
        "products": product_data,
        "matrix": matrix,
        "display_names": display_map,
    }


@router.get("/products/export")
def export_products(
    category_id: Optional[int] = None, search: str = "",
    db: Session = Depends(get_db), user=Depends(get_current_user),
):
    q = db.query(Product).options(
        selectinload(Product.category),
        selectinload(Product.manufacturer),
        selectinload(Product.supplier),
        selectinload(Product.comm_methods).selectinload(ProductCommMethod.method),
        selectinload(Product.comm_protocols).selectinload(ProductCommProtocol.protocol),
        selectinload(Product.power_supplies).selectinload(ProductPowerSupply.power),
    )
    if category_id:
        from app.services.product_category_helper import get_products_in_categories, get_category_descendants
        cat_ids = get_category_descendants(db, category_id)
        m2m_ids = get_products_in_categories(db, cat_ids)
        q = q.filter(
            or_(
                Product.id.in_(m2m_ids),
                Product.category_id.in_(cat_ids)
            )
        )
    if search:
        q = q.filter(or_(
            Product.name.ilike(f"%{escape_like(search)}%"),
            Product.model.ilike(f"%{escape_like(search)}%"),
            Product.description.ilike(f"%{escape_like(search)}%"),
            Product.id.in_(
                select(text('pc.product_id')).select_from(text('product_categories pc JOIN device_categories c ON pc.category_id = c.id')).where(text('c.name LIKE :s')).params(s=f'%{escape_like(search)}%')
            ),
        ))
    _mfg_pri_ids2 = [m.id for m in db.query(Manufacturer).filter(Manufacturer.sort_order < 100).all()]
    _sk2 = case((Product.manufacturer_id.in_(_mfg_pri_ids2), 0), (Product.image_url != None, 1), else_=2)
    products = q.order_by(_sk2, Product.name).all()

    from app.utils.excel_style import (
        apply_info_row, apply_title_row, apply_header_row, apply_data_row,
        apply_footer_row, apply_column_widths,
    )
    from datetime import date

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "产品清单"
    apply_column_widths(ws)

    # Row 1: info
    apply_info_row(ws, 1, f"导出人：{user.username}  |  日期：{date.today().isoformat()}  |  共 {len(products)} 条")

    # Row 2: title
    apply_title_row(ws, 2, "产品清单")

    # Row 3: headers
    headers = ["序号", "名称", "规格型号", "型号", "功能描述", "单价", "品类", "厂商", "通讯", "供电", "备注", "图片"]
    apply_header_row(ws, 3, headers)

    # Data rows
    for idx, p in enumerate(products, 1):
        comm = ", ".join(cm.method.name for cm in p.comm_methods if cm.method)
        proto = ", ".join(cp.protocol.name for cp in p.comm_protocols if cp.protocol)
        power = ", ".join(ps.power.name for ps in p.power_supplies if ps.power)
        all_comm = ", ".join(filter(None, [comm, proto]))
        cat_name = p.category.name if p.category else ""
        apply_data_row(ws, 3 + idx, [
            idx,
            p.name,
            p.model or "",
            p.model or "",
            format_description_with_specs(p.description or "", p.specs or {}),
            float(p.base_price or 0),
            cat_name,
            p.manufacturer.name if p.manufacturer else "",
            all_comm,
            power,
            "",
            p.image_url or "",
        ])

    # Footer
    footer_row = 3 + len(products) + 1
    apply_footer_row(ws, footer_row, "产品数据库 — 产品清单导出")

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename*=UTF-8''products.xlsx"},
    )


@router.post("/products/upload-image")
def upload_image(file: UploadFile = File(...), user=Depends(get_current_user)):
    """Upload an image file, return URL."""
    if not file.content_type or not file.content_type.startswith("image/"):
        raise HTTPException(400, "File must be an image")
    # Check size before reading to avoid memory exhaustion
    if file.size and file.size > settings.IMAGE_MAX_SIZE:
        raise HTTPException(400, "Image must be under 5MB")
    content = file.file.read()
    if len(content) > settings.IMAGE_MAX_SIZE:
        raise HTTPException(400, "Image must be under 5MB")

    from app.services.storage import ALLOWED_EXTENSIONS, VALID_SIGNATURES
    # Validate file extension
    original_name = file.filename or "upload.jpg"
    ext = os.path.splitext(original_name)[1].lower()
    if ext not in {e for e in ALLOWED_EXTENSIONS if e.startswith('.') and e in ('.jpg','.jpeg','.png','.gif','.webp','.bmp')}:
        raise HTTPException(400, f"File extension '{ext}' not allowed. Allowed: .jpg, .jpeg, .png, .gif, .webp, .bmp")

    # Magic bytes validation (shared constants from storage.py)
    if len(content) >= 8:
        header = content[:8]
        if not any(header.startswith(sig) for sig in VALID_SIGNATURES):
            raise HTTPException(400, "File content does not match a valid image format")
    else:
        raise HTTPException(400, "File too small to be a valid image")

    # Use UUID filename to avoid path traversal
    safe_name = f"{uuid.uuid4().hex[:16]}{ext}"
    url = save_upload(content, safe_name)
    return {"url": url}


@router.post("/products/download-image")
def download_image(data: DownloadImageRequest, user=Depends(get_current_user)):
    """Download image from URL, save locally, return URL."""
    url = data.url.strip()
    if not url:
        raise HTTPException(400, "URL is required")
    if not validate_url(url):
        raise HTTPException(400, "URL not allowed")
    try:
        local_url = upload_from_url(url)
        return {"url": local_url}
    except Exception as e:
        raise HTTPException(400, f"Failed to download image: {e}")


# --- Parameterized routes ---

@router.get("/products/{product_id}")
def get_product(product_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    p = db.scalar(select(Product).options(*product_eager_loads()).where(Product.id == product_id))
    if not p:
        raise HTTPException(404, "Product not found")
    p.view_count = (p.view_count or 0) + 1
    db.commit()
    result = build_product_detail(p, db)
    from app.services.field_visibility import filter_fields_for_user
    is_admin = getattr(user, 'role', '') == 'admin'
    filter_fields_for_user(result, is_admin, db)
    return {"product": result}


@router.post("/products")
def create_product(data: ProductCreate, db: Session = Depends(get_db), user=Depends(get_current_user)):
    # Validate specs if category has definitions
    if data.specs and data.category_id:
        spec_defs = db.query(CategorySpecDefinition).filter_by(category_id=data.category_id).all()
        specs = data.specs
        errors = validate_specs(specs, spec_defs)
        if errors:
            raise HTTPException(400, detail={"errors": errors})

    p = Product(
        model=data.model,
        name=data.name,
        sku=data.sku,
        category_id=data.category_id,
        manufacturer_id=data.manufacturer_id,
        supplier_id=data.supplier_id,
        unit=data.unit,
        base_price=data.base_price,
        cost_price=data.cost_price,
        description=data.description,
        image_url=data.image_url,
        product_url=data.product_url,
        status=data.status,
        parent_id=data.parent_id,
        specs=data.specs,
        urls=data.urls,
        custom_fields=data.custom_fields,
        created_by=user.id,
    )
    p.pinyin_search = build_pinyin(f"{p.name} {p.model or ''}")
    p.created_at = datetime.now()
    p.updated_at = datetime.now()

    db.add(p)
    db.flush()  # get p.id

    # Write multi-category relationships
    from app.services.product_category_helper import add_product_categories
    add_product_categories(db, p.id, data.category_ids or [data.category_id])

    # Write mapping tables
    write_mappings(p.id, data.model_dump(exclude_none=True), db)

    db.commit()
    db.refresh(p)
    cats, mfgs, sups = get_name_maps(db)
    result = p.to_dict(cats, sups, mfgs)
    from app.services.field_visibility import filter_fields_for_user
    is_admin = getattr(user, 'role', '') == 'admin'
    filter_fields_for_user(result, is_admin, db)
    return {"product": result}


@router.put("/products/{product_id}")
def update_product(product_id: int, data: ProductUpdate, db: Session = Depends(get_db), user=Depends(get_current_user)):
    p = get_or_404(db, Product, product_id)
    check_ownership(p, user)

    from app.utils.helpers import apply_partial_update
    apply_partial_update(p, data, ["model", "name", "sku", "category_id", "manufacturer_id",
        "supplier_id", "unit", "base_price", "cost_price", "description", "image_url",
        "product_url", "status", "parent_id"])

    if data.specs:
        p.specs = data.specs
    if data.urls:
        p.urls = data.urls
    if data.custom_fields:
        p.custom_fields = data.custom_fields

    p.pinyin_search = build_pinyin(f"{p.name} {p.model or ''}")
    p.updated_at = datetime.now()

    # Update multi-category if provided
    if data.category_ids is not None:
        if not data.category_ids or not any(data.category_ids):
            raise HTTPException(400, '品类不能为空')
        from app.services.product_category_helper import add_product_categories
        add_product_categories(db, product_id, data.category_ids)
        p.category_id = data.category_ids[0]

    # Rewrite mapping tables (delete old, insert new) — only for fields in payload
    payload = data.model_dump(exclude_none=True)
    rewrite_mappings(product_id, payload, db)

    db.commit()
    cats, mfgs, sups = get_name_maps(db)
    result = p.to_dict(cats, sups, mfgs)
    from app.services.field_visibility import filter_fields_for_user
    is_admin = getattr(user, 'role', '') == 'admin'
    filter_fields_for_user(result, is_admin, db)
    return {"product": result}


@router.delete("/products/{product_id}")
def delete_product(product_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    p = get_or_404(db, Product, product_id)
    check_ownership(p, user)
    # Delete local image files before removing DB rows
    from app.services.product_helpers import _cleanup_image_files
    _cleanup_image_files(product_id, db)
    # Delete dependencies before product (SQLite FK cascade not enabled)
    from app.models.dependency import ProductDependency
    db.query(ProductDependency).filter(
        (ProductDependency.product_id == product_id) | (ProductDependency.depends_on_product_id == product_id)
    ).delete()
    db.delete(p)
    db.commit()
    return {"ok": True}


# --- AI-assisted URL fetch ---


@router.post("/products/ai-fetch")
def ai_fetch_product(data: AIFetchRequest, db: Session = Depends(get_db), user=Depends(get_current_user)):
    """AI-assisted product info extraction from URL or text."""
    url = data.url.strip()
    raw_text = data.text.strip()

    if raw_text:
        # Direct text extraction mode
        from app.services.ai_engine import engine
        usage: list = []
        result = call_ai_extract("", "", raw_text, db, usage) if engine.api_key \
            else regex_extract_from_text("", raw_text, db)
        if usage:
            _log_ai_usage(user.id, 'ai_fetch_text', usage[0], usage[1], usage[2])
        elif engine.api_key:
            _log_ai_usage(user.id, 'ai_fetch_text')
        return {"fetched": result}

    if not url:
        raise HTTPException(400, "URL or text is required")

    parsed = urlparse(url)
    if not parsed.scheme or not parsed.netloc:
        raise HTTPException(400, "Invalid URL")

    if not validate_url(url):
        raise HTTPException(400, "URL not allowed")

    try:
        resp = httpx.get(url, timeout=30, follow_redirects=False, headers={
            "User-Agent": "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36"
        })
        # Manually follow redirects with SSRF validation
        for _ in range(3):
            if resp.status_code in (301, 302, 303, 307, 308):
                loc = resp.headers.get("Location", "")
                if not loc or not validate_url(loc):
                    raise HTTPException(400, f"Redirect target not allowed: {loc[:60]}")
                resp = httpx.get(loc, timeout=30, follow_redirects=False, headers={
                    "User-Agent": "Mozilla/5.0"
                })
            else:
                break
        resp.raise_for_status()
    except httpx.HTTPError as e:
        raise HTTPException(400, f"Failed to fetch URL: {str(e)}")

    soup = BeautifulSoup(resp.text, "html.parser")
    for tag in soup(["script", "style", "nav", "footer", "header", "noscript", "svg", "path"]):
        tag.decompose()

    title = soup.title.string.strip() if soup.title else ""
    body = soup.body
    text = body.get_text(separator="\n", strip=True) if body else resp.text
    text = re.sub(r'\n{3,}', '\n\n', text)[:12000]

    if not text:
        raise HTTPException(400, "No text content extracted from URL")

    usage: list = []
    result = call_ai_extract(url, title, text, db, usage)
    if usage:
        _log_ai_usage(user.id, 'ai_fetch', usage[0], usage[1], usage[2])
    else:
        _log_ai_usage(user.id, 'ai_fetch')

    return {"fetched": {"url": url, "title": title, **result}}


def _log_ai_usage(user_id: int, operation: str, tokens_in: int = 0, tokens_out: int = 0, duration_ms: int = 0, success: bool = True):
    """Record AI usage for stats sidebar. Uses lazy import to avoid circular deps."""
    _logger = logging.getLogger("uvicorn")
    try:
        from app.models.ai_usage_log import AIUsageLog
        from app.database import SessionLocal
        db2 = SessionLocal()
        try:
            db2.add(AIUsageLog(user_id=user_id, operation=operation,
                tokens_in=tokens_in, tokens_out=tokens_out,
                duration_ms=duration_ms, success=success))
            db2.commit()
        finally:
            db2.close()
    except Exception:
        _logger.warning("Failed to log AI usage for user %d, op %s", user_id, operation)


@router.post("/products/ai-fetch-file")
async def ai_fetch_file(file: UploadFile = File(...), db: Session = Depends(get_db), user=Depends(get_current_user)):
    """AI-assisted product info extraction from uploaded file (PDF/docx/txt/image)."""
    content = await file.read()
    filename = (file.filename or "").lower()
    text = ""
    is_image = False

    if filename.endswith(".pdf"):
        import io
        from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(content))
        text = "\n".join(p.extract_text() or "" for p in reader.pages)
    elif filename.endswith(".docx"):
        import io
        from docx import Document
        doc = Document(io.BytesIO(content))
        text = "\n".join(p.text for p in doc.paragraphs)
    elif filename.endswith(".txt") or filename.endswith(".csv"):
        text = content.decode("utf-8", errors="replace")
    elif filename.endswith((".jpg", ".jpeg", ".png", ".webp", ".bmp", ".gif")):
        is_image = True
    else:
        raise HTTPException(400, f"Unsupported file type: {filename.split('.')[-1]}")

    # Image: OCR via vision model first
    ocr_usage: list = []
    if is_image:
        text = await _ocr_image(content, db, ocr_usage)

    text = re.sub(r'\n{3,}', '\n\n', (text or "")[:8000])
    if not text.strip():
        raise HTTPException(400, "No text content extracted from file")

    usage: list = []
    result = await _extract_product_info(text, filename, db, usage)
    if usage:
        _log_ai_usage(user.id, 'ai_fetch_file', usage[0], usage[1], usage[2])
    if ocr_usage:
        _log_ai_usage(user.id, 'ai_ocr', ocr_usage[0], ocr_usage[1], ocr_usage[2])
    return {"fetched": {"filename": filename, **result}}


async def _ocr_image(image_bytes: bytes, db: Session, _usage: list = None) -> str:
    """Send image to vision LLM for OCR text extraction.
    If _usage list is provided, [tokens_in, tokens_out, duration_ms] is appended.
    """
    import base64, time as _time
    from app.routers.admin_routes import _load_llm_config

    cfg = _load_llm_config(db).get("vision", {})
    if not cfg.get("api_key"):
        raise HTTPException(400, "Vision LLM API key not configured")

    mime_map = {b'\xff\xd8\xff': 'image/jpeg', b'\x89PNG': 'image/png',
                b'GIF8': 'image/gif', b'RIFF': 'image/webp', b'BM': 'image/bmp'}
    mime_type = "image/jpeg"
    for sig, mt in mime_map.items():
        if image_bytes[:len(sig)] == sig:
            mime_type = mt; break

    b64 = base64.b64encode(image_bytes).decode()
    data_url = f"data:{mime_type};base64,{b64}"

    import httpx
    t0 = _time.time()
    try:
        async with httpx.AsyncClient(timeout=60) as client:
            resp = await client.post(
                f"{cfg['base_url']}/chat/completions",
                headers={"Authorization": f"Bearer {cfg['api_key']}", "Content-Type": "application/json"},
                json={
                    "model": cfg.get("model", "mimo-v2-omni"),
                    "messages": [{"role": "user", "content": [
                        {"type": "text", "text": "请仔细OCR识别这张产品规格图片中的所有文字内容，包括产品名称、型号、规格参数、技术指标等。直接输出识别到的文字，不要添加额外说明。"},
                        {"type": "image_url", "image_url": {"url": data_url}},
                    ]}],
                    "max_tokens": 2000,
                },
            )
        resp.raise_for_status()
        data = resp.json()
        if _usage is not None:
            u = data.get("usage", {})
            _usage.append(u.get("prompt_tokens", 0))
            _usage.append(u.get("completion_tokens", 0))
            _usage.append(int((_time.time() - t0) * 1000))
        return data["choices"][0]["message"]["content"] or ""
    except Exception:
        raise HTTPException(500, "Vision LLM OCR failed")


async def _extract_product_info(text: str, source: str, db: Session, _usage: list = None) -> dict:
    """Extract product info from text using AI or regex fallback.
    If _usage list is provided, [tokens_in, tokens_out, duration_ms] is appended on success.
    """
    from app.services.ai_engine import engine
    if engine.api_key:
        from app.services.ai_extract import build_extraction_prompt
        system_prompt = build_extraction_prompt(db)
        user_msg = f"来源: {source}\n\n规格文本:\n{text}\n\n请提取产品的结构化信息。"
        import re as _re, time as _time
        try:
            t0 = _time.time()
            result = await engine.chat([
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_msg},
            ], temperature=0.1, max_tokens=3000)
            elapsed_ms = int((_time.time() - t0) * 1000)
            u = result.get("usage", {})
            if _usage is not None:
                _usage.append(u.get("prompt_tokens", 0))
                _usage.append(u.get("completion_tokens", 0))
                _usage.append(elapsed_ms)
            content_text = result["choices"][0]["message"]["content"]
            json_match = _re.search(r'\{.*\}', content_text, re.DOTALL)
            if json_match:
                return json.loads(json_match.group())
        except Exception:
            pass
    return regex_extract_from_text("", text, db)




@router.get("/products/{product_id}/spec-sheet")
def spec_sheet(product_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    """Generate product spec sheet as PDF download or HTML view."""
    p = db.scalar(select(Product).options(*product_eager_loads()).where(Product.id == product_id))
    if not p:
        raise HTTPException(404, "Product not found")

    html = build_spec_html(p, db)
    filename = re.sub(r'[^\w\-_.]', '_', p.name or 'product')

    # Try CLI weasyprint for PDF
    import shutil
    wp = settings.WEASYPRINT_PATH or shutil.which("weasyprint")
    if wp:
        try:
            with tempfile.NamedTemporaryFile(mode="w", suffix=".html", delete=False) as html_f:
                html_f.write(html)
                html_path = html_f.name
            pdf_fd, pdf_path = tempfile.mkstemp(suffix=".pdf")
            os.close(pdf_fd)
            env = os.environ.copy()
            import sys
            if sys.platform == "darwin":
                env["DYLD_LIBRARY_PATH"] = "/opt/homebrew/lib"
            elif sys.platform == "linux":
                env["LD_LIBRARY_PATH"] = env.get("LD_LIBRARY_PATH", "")
            result = subprocess.run(
                [wp, html_path, pdf_path, "-e", "utf-8"],
                env=env, capture_output=True,
            )
            if result.returncode == 0 and os.path.exists(pdf_path) and os.path.getsize(pdf_path) > 100:
                with open(pdf_path, "rb") as f2:
                    pdf_bytes = f2.read()
                os.unlink(html_path)
                os.unlink(pdf_path)
                return StreamingResponse(
                    BytesIO(pdf_bytes),
                    media_type="application/pdf",
                    headers={"Content-Disposition": f"attachment; filename={filename}-spec-sheet.pdf"},
                )
        except Exception as e:
            logging.getLogger("uvicorn").warning(f"PDF generation failed for product {product_id}: {e}")
        finally:
            for path in [html_path, pdf_path]:
                try:
                    if os.path.exists(path): os.unlink(path)
                except OSError:
                    pass

    return HTMLResponse(content=html)


# --- Dependencies ---

@router.get("/products/{product_id}/dependencies")
def get_dependencies(product_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    deps = db.query(ProductDependency).filter_by(product_id=product_id).all()
    return {"dependencies": [d.to_dict() for d in deps]}


@router.post("/products/{product_id}/dependencies")
def create_dependency(product_id: int, data: ProductDependencyCreate, db: Session = Depends(get_db), user=Depends(get_current_user)):
    dep = ProductDependency(
        product_id=product_id,
        depends_on_product_id=data.depends_on_product_id,
        depends_on_category_id=data.depends_on_category_id,
        dependency_type=data.dependency_type,
        description=data.description,
        sort_order=data.sort_order,
    )
    db.add(dep)
    db.commit()
    db.refresh(dep)
    return {"dependency": dep.to_dict()}


@router.put("/products/{product_id}/dependencies/{dep_id}")
def update_dependency(product_id: int, dep_id: int, data: ProductDependencyUpdate, db: Session = Depends(get_db), user=Depends(get_current_user)):
    dep = db.get(ProductDependency, dep_id)
    if not dep or dep.product_id != product_id:
        raise HTTPException(404)
    apply_partial_update(dep, data, ["depends_on_product_id", "depends_on_category_id", "dependency_type", "description", "sort_order"])
    db.commit()
    return {"dependency": dep.to_dict()}


@router.delete("/products/{product_id}/dependencies/{dep_id}")
def delete_dependency(product_id: int, dep_id: int, db: Session = Depends(get_db), user=Depends(get_current_user)):
    dep = db.get(ProductDependency, dep_id)
    if not dep or dep.product_id != product_id:
        raise HTTPException(404)
    db.delete(dep)
    db.commit()
    return {"ok": True}
