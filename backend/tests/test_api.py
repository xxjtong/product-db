"""Comprehensive test suite for product-db — uses a separate test database."""
import pytest
import os
import json
import tempfile

# Force test DB before any app imports
_test_db_path = tempfile.mktemp(suffix=".db")
os.environ["DATABASE_URL"] = f"sqlite:///{_test_db_path}"
os.environ["DEV_MODE"] = "true"
os.environ["SECRET_KEY"] = "test-secret-key-for-pytest-32charmin"

from fastapi.testclient import TestClient
from app.database import Base, engine, SessionLocal, get_db
from app.main import app
from app.models.user import User
from app.models.category import Category, CategorySpecDefinition
from app.models.product import Product
from app.models.supplier import Supplier
from app.models.dictionary import Manufacturer, DictCommMethod, DictCommProtocol, DictPowerSupply
from app.models.dependency import ProductDependency
from app.models.solution import Solution, SolutionItem
from app.models.quotation import Quotation, QuotationItem
from app.models.bom_template import BOMTemplate, SolutionBOMSnapshot
from app.auth import hash_password, create_token
from tests.conftest import create_test_schema, drop_test_schema

client = TestClient(app)


@pytest.fixture(autouse=True)
def setup_db():
    """Create all tables and seed admin user before each test."""
    create_test_schema()
    db = SessionLocal()
    if not db.query(User).filter_by(username="admin").first():
        db.add(User(username="admin", password_hash=hash_password("admin"), role="admin"))
        db.commit()
    db.close()
    yield
    drop_test_schema()


@pytest.fixture
def db():
    """Provide a DB session for direct data setup."""
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


@pytest.fixture
def auth_headers(db):
    """Return auth headers with a valid JWT."""
    admin = db.query(User).filter_by(username="admin").first()
    token = create_token(admin.id, admin.username)
    return {"Authorization": f"Bearer {token}"}


def _seed_category(db, name="测试品类", slug="test-cat"):
    cat = Category(name=name, slug=slug, level=1, sort_order=0)
    db.add(cat)
    db.commit()
    db.refresh(cat)
    return cat


def _seed_manufacturer(db, name="测试厂商"):
    mfg = Manufacturer(name=name)
    db.add(mfg)
    db.commit()
    db.refresh(mfg)
    return mfg


def _seed_supplier(db, name="测试供应商"):
    s = Supplier(name=name)
    db.add(s)
    db.commit()
    db.refresh(s)
    return s


def _seed_product(db, name="测试产品", model="TEST-001", category_id=1, manufacturer_id=None, supplier_id=None, **kwargs):
    init_kwargs = dict(kwargs)
    p = Product(name=name, model=model, category_id=category_id,
                manufacturer_id=manufacturer_id, supplier_id=supplier_id, **init_kwargs)
    db.add(p)
    db.commit()
    db.refresh(p)
    return p


# ============================================================
# Health
# ============================================================
class TestHealth:
    def test_health(self):
        resp = client.get("/product-db/api/health")
        assert resp.status_code == 200
        assert resp.json() == {"status": "ok"}


# ============================================================
# Auth
# ============================================================
class TestAuth:
    def test_login_fail(self):
        resp = client.post("/product-db/api/auth/login", json={"username": "nobody", "password": "wrong"})
        assert resp.status_code == 401

    def test_login_success(self):
        resp = client.post("/product-db/api/auth/login", json={"username": "admin", "password": "admin"})
        assert resp.status_code in (200, 201)
        data = resp.json()
        assert "token" in data

    def test_me(self):
        resp = client.get("/product-db/api/auth/me")
        assert resp.status_code == 200
        assert "user" in resp.json()

    def test_admin_users(self):
        resp = client.get("/product-db/api/admin/users")
        assert resp.status_code == 200


# ============================================================
# Categories
# ============================================================
class TestCategories:
    def test_create_and_list(self, db):
        cat = _seed_category(db)
        resp = client.get("/product-db/api/categories")
        assert resp.status_code == 200
        data = resp.json()
        assert data["total"] >= 1
        assert any(c["name"] == "测试品类" for c in data["categories"])

    def test_tree(self, db):
        _seed_category(db, name="父品类", slug="parent")
        resp = client.get("/product-db/api/categories/tree")
        assert resp.status_code == 200
        assert "tree" in resp.json()

    def test_create_category(self, db):
        resp = client.post("/product-db/api/categories", json={"name": "新建品类", "slug": "new-cat"})
        assert resp.status_code in (200, 201)
        assert resp.json()["category"]["name"] == "新建品类"

    def test_update_category(self, db):
        cat = _seed_category(db)
        resp = client.put(f"/product-db/api/categories/{cat.id}", json={"name": "更新品类"})
        assert resp.status_code == 200
        assert resp.json()["category"]["name"] == "更新品类"

    def test_delete_category(self, db):
        cat = _seed_category(db)
        cat_id = cat.id
        resp = client.delete(f"/product-db/api/categories/{cat_id}")
        assert resp.status_code == 200
        db.expire_all()
        assert db.get(Category, cat_id) is None

    def test_spec_definitions_crud(self, db):
        cat = _seed_category(db)
        # Create
        resp = client.post(f"/product-db/api/categories/{cat.id}/spec-definitions", json={
            "spec_key": "ip_rating", "display_name": "IP等级",
            "spec_type": "enum", "options": ["IP30", "IP65", "IP67"],
        })
        assert resp.status_code in (200, 201)
        sd_id = resp.json()["spec_definition"]["id"]

        # List
        resp = client.get(f"/product-db/api/categories/{cat.id}/spec-definitions")
        assert resp.status_code == 200
        assert len(resp.json()["spec_definitions"]) == 1

        # Update
        resp = client.put(f"/product-db/api/categories/{cat.id}/spec-definitions/{sd_id}",
                          json={"display_name": "防护等级"})
        assert resp.status_code == 200
        assert resp.json()["spec_definition"]["display_name"] == "防护等级"

        # Delete
        resp = client.delete(f"/product-db/api/categories/{cat.id}/spec-definitions/{sd_id}")
        assert resp.status_code == 200


# ============================================================
# Suppliers
# ============================================================
class TestSuppliers:
    def test_create_and_list(self, db):
        _seed_supplier(db)
        resp = client.get("/product-db/api/suppliers")
        assert resp.status_code == 200
        assert resp.json()["total"] >= 1

    def test_search(self, db):
        _seed_supplier(db, name="深圳科技")
        resp = client.get("/product-db/api/suppliers?search=深圳")
        assert resp.status_code == 200
        assert any(s["name"] == "深圳科技" for s in resp.json()["suppliers"])

    def test_update(self, db):
        s = _seed_supplier(db)
        resp = client.put(f"/product-db/api/suppliers/{s.id}", json={"name": "更新供应商"})
        assert resp.status_code == 200
        assert resp.json()["supplier"]["name"] == "更新供应商"

    def test_delete(self, db):
        s = _seed_supplier(db)
        s_id = s.id
        resp = client.delete(f"/product-db/api/suppliers/{s_id}")
        assert resp.status_code == 200
        db.expire_all()
        assert db.get(Supplier, s_id) is None

    def test_404_update(self):
        resp = client.put("/product-db/api/suppliers/99999", json={"name": "x"})
        assert resp.status_code == 404


# ============================================================
# Products
# ============================================================
class TestProducts:
    def test_create_and_list(self, db):
        cat = _seed_category(db)
        _seed_product(db, category_id=cat.id)
        resp = client.get("/product-db/api/products")
        assert resp.status_code == 200
        assert resp.json()["total"] >= 1

    def test_filter_by_category(self, db):
        cat = _seed_category(db)
        _seed_product(db, category_id=cat.id)
        resp = client.get(f"/product-db/api/products?category_id={cat.id}")
        assert resp.status_code == 200
        assert resp.json()["total"] >= 1

    def test_search(self, db):
        cat = _seed_category(db)
        _seed_product(db, name="温度传感器", category_id=cat.id)
        resp = client.get("/product-db/api/products?search=温度")
        assert resp.status_code == 200

    def test_detail(self, db):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)
        resp = client.get(f"/product-db/api/products/{p.id}")
        assert resp.status_code == 200
        assert resp.json()["product"]["name"] == "测试产品"

    def test_detail_404(self):
        resp = client.get("/product-db/api/products/99999")
        assert resp.status_code == 404

    def test_create_product(self, db):
        cat = _seed_category(db)
        mfg = _seed_manufacturer(db)
        resp = client.post("/product-db/api/products", json={
            "name": "新产品", "model": "NEW-01",
            "category_id": cat.id, "manufacturer_id": mfg.id,
            "specs": {"ip_rating": "IP65"},
        })
        assert resp.status_code in (200, 201)
        assert resp.json()["product"]["name"] == "新产品"

    def test_update_product(self, db):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)
        resp = client.put(f"/product-db/api/products/{p.id}", json={"name": "更新产品"})
        assert resp.status_code == 200
        assert resp.json()["product"]["name"] == "更新产品"

    def test_delete_product(self, db):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)
        p_id = p.id
        resp = client.delete(f"/product-db/api/products/{p_id}")
        assert resp.status_code == 200
        db.expire_all()
        assert db.get(Product, p_id) is None

    def test_compare(self, db):
        cat = _seed_category(db)
        p1 = _seed_product(db, name="产品A", model="A", category_id=cat.id, specs={"ip_rating": "IP65"})
        p2 = _seed_product(db, name="产品B", model="B", category_id=cat.id, specs={"ip_rating": "IP67"})
        resp = client.get(f"/product-db/api/products/compare?product_ids={p1.id},{p2.id}")
        assert resp.status_code == 200
        assert "matrix" in resp.json()

    def test_compare_min_ids(self):
        resp = client.get("/product-db/api/products/compare?product_ids=1")
        assert resp.status_code == 400

    def test_spec_sheet(self, db):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)
        resp = client.get(f"/product-db/api/products/{p.id}/spec-sheet")
        assert resp.status_code == 200

    def test_export(self, db):
        cat = _seed_category(db)
        _seed_product(db, category_id=cat.id)
        resp = client.get("/product-db/api/products/export")
        assert resp.status_code == 200
        assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")

    def test_ai_fetch_text(self):
        resp = client.post("/product-db/api/products/ai-fetch", json={"text": "UG67 LoRaWAN Gateway IP67 Milesight"})
        assert resp.status_code in (200, 201)
        assert "fetched" in resp.json()

    def test_ai_fetch_empty(self):
        resp = client.post("/product-db/api/products/ai-fetch", json={"url": "", "text": ""})
        assert resp.status_code == 400


# ============================================================
# Product Dependencies
# ============================================================
class TestDependencies:
    def test_create_and_list(self, db):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)
        target_cat = _seed_category(db, name="网关", slug="gateway")
        resp = client.post(f"/product-db/api/products/{p.id}/dependencies", json={
            "depends_on_category_id": target_cat.id,
            "dependency_type": "required",
            "description": "需要网关",
        })
        assert resp.status_code in (200, 201)
        dep_id = resp.json()["dependency"]["id"]

        resp = client.get(f"/product-db/api/products/{p.id}/dependencies")
        assert resp.status_code == 200
        assert len(resp.json()["dependencies"]) == 1

    def test_update_dependency(self, db):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)
        target_cat = _seed_category(db, name="网关", slug="gateway")
        dep = ProductDependency(product_id=p.id, depends_on_category_id=target_cat.id,
                                dependency_type="required")
        db.add(dep)
        db.commit()
        db.refresh(dep)

        resp = client.put(f"/product-db/api/products/{p.id}/dependencies/{dep.id}",
                          json={"dependency_type": "recommended"})
        assert resp.status_code == 200
        assert resp.json()["dependency"]["dependency_type"] == "recommended"

    def test_delete_dependency(self, db):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)
        target_cat = _seed_category(db, name="网关", slug="gateway")
        dep = ProductDependency(product_id=p.id, depends_on_category_id=target_cat.id,
                                dependency_type="required")
        db.add(dep)
        db.commit()
        db.refresh(dep)
        dep_id = dep.id

        resp = client.delete(f"/product-db/api/products/{p.id}/dependencies/{dep_id}")
        assert resp.status_code == 200
        db.expire_all()
        assert db.get(ProductDependency, dep_id) is None


# ============================================================
# Solutions
# ============================================================
class TestSolutions:
    def test_create_and_list(self, db):
        resp = client.post("/product-db/api/solutions", json={"name": "测试方案"})
        assert resp.status_code in (200, 201)
        sol_id = resp.json()["solution"]["id"]

        resp = client.get("/product-db/api/solutions")
        assert resp.status_code == 200
        assert resp.json()["total"] >= 1

    def test_get_solution(self, db):
        resp = client.post("/product-db/api/solutions", json={"name": "方案详情"})
        sol_id = resp.json()["solution"]["id"]

        resp = client.get(f"/product-db/api/solutions/{sol_id}")
        assert resp.status_code in (200, 201)
        assert resp.json()["solution"]["name"] == "方案详情"

    def test_update_solution(self, db):
        resp = client.post("/product-db/api/solutions", json={"name": "原始方案"})
        sol_id = resp.json()["solution"]["id"]

        resp = client.put(f"/product-db/api/solutions/{sol_id}", json={"name": "更新方案"})
        assert resp.status_code in (200, 201)
        assert resp.json()["solution"]["name"] == "更新方案"

    def test_delete_solution(self, db):
        resp = client.post("/product-db/api/solutions", json={"name": "删除方案"})
        sol_id = resp.json()["solution"]["id"]

        resp = client.delete(f"/product-db/api/solutions/{sol_id}")
        assert resp.status_code in (200, 201)

    def test_solution_items_crud(self, db):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)

        # Create solution
        resp = client.post("/product-db/api/solutions", json={"name": "BOM方案"})
        sol_id = resp.json()["solution"]["id"]

        # Add item
        resp = client.post(f"/product-db/api/solutions/{sol_id}/items", json={
            "product_id": p.id, "quantity": 2, "unit_price": 1500,
        })
        assert resp.status_code in (200, 201)
        item_id = resp.json()["item"]["id"]

        # List items
        resp = client.get(f"/product-db/api/solutions/{sol_id}/items")
        assert resp.status_code == 200
        assert len(resp.json()["items"]) == 1

        # Update item
        resp = client.put(f"/product-db/api/solutions/{sol_id}/items/{item_id}",
                          json={"quantity": 5})
        assert resp.status_code == 200
        assert resp.json()["item"]["quantity"] == 5

        # Delete item
        resp = client.delete(f"/product-db/api/solutions/{sol_id}/items/{item_id}")
        assert resp.status_code == 200


# ============================================================
# Solution Check & Suggest
# ============================================================
class TestSolutionCheck:
    def test_check_no_warnings(self, db):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)

        resp = client.post("/product-db/api/solutions", json={"name": "完整方案"})
        sol_id = resp.json()["solution"]["id"]

        client.post(f"/product-db/api/solutions/{sol_id}/items", json={
            "product_id": p.id, "quantity": 1,
        })

        resp = client.get(f"/product-db/api/solutions/{sol_id}/check")
        assert resp.status_code in (200, 201)
        assert resp.json()["ok"] is True
        assert len(resp.json()["warnings"]) == 0

    def test_check_missing_dependency(self, db):
        cat = _seed_category(db)
        gateway_cat = _seed_category(db, name="网关", slug="gateway")
        sensor = _seed_product(db, name="传感器", category_id=cat.id)

        # Add required dependency: sensor needs a gateway
        dep = ProductDependency(product_id=sensor.id, depends_on_category_id=gateway_cat.id,
                                dependency_type="required", description="需要LoRaWAN网关")
        db.add(dep)
        db.commit()

        # Create solution with only sensor (no gateway)
        resp = client.post("/product-db/api/solutions", json={"name": "缺网关方案"})
        sol_id = resp.json()["solution"]["id"]
        client.post(f"/product-db/api/solutions/{sol_id}/items", json={
            "product_id": sensor.id, "quantity": 10,
        })

        resp = client.get(f"/product-db/api/solutions/{sol_id}/check")
        assert resp.status_code in (200, 201)
        assert resp.json()["ok"] is False
        assert len(resp.json()["warnings"]) >= 1
        assert resp.json()["warnings"][0]["type"] == "missing_category"

    def test_suggest_fulfills_dependency(self, db):
        cat = _seed_category(db)
        gateway_cat = _seed_category(db, name="网关", slug="gateway")
        sensor = _seed_product(db, name="传感器", category_id=cat.id)
        gateway = _seed_product(db, name="UG65网关", model="UG65", category_id=gateway_cat.id,
                                specs={"max_endpoints": 500})

        dep = ProductDependency(product_id=sensor.id, depends_on_category_id=gateway_cat.id,
                                dependency_type="required")
        db.add(dep)
        db.commit()

        resp = client.post("/product-db/api/solutions", json={"name": "需推荐方案"})
        sol_id = resp.json()["solution"]["id"]
        client.post(f"/product-db/api/solutions/{sol_id}/items", json={
            "product_id": sensor.id, "quantity": 5,
        })

        resp = client.get(f"/product-db/api/solutions/{sol_id}/suggest")
        assert resp.status_code in (200, 201)
        suggestions = resp.json()["suggestions"]
        assert len(suggestions) >= 1
        assert suggestions[0]["missing_category"] == "网关"
        assert any(prod["name"] == "UG65网关" for prod in suggestions[0]["products"])


# ============================================================
# Quotations
# ============================================================
class TestQuotations:
    def test_create_and_list(self, db):
        resp = client.post("/product-db/api/quotations", json={"title": "测试报价单"})
        assert resp.status_code in (200, 201)
        qt = resp.json()["quotation"]
        assert qt["title"] == "测试报价单"
        assert qt["quote_number"].startswith("QT-")

        resp = client.get("/product-db/api/quotations")
        assert resp.status_code == 200
        assert resp.json()["total"] >= 1

    def test_create_from_solution(self, db):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)
        p.base_price = 1500
        db.commit()

        sol = Solution(name="方案A", client_name="客户甲")
        db.add(sol)
        db.commit()
        db.refresh(sol)

        item = SolutionItem(solution_id=sol.id, product_id=p.id, quantity=2,
                            unit_price=1500, discount_rate=100)
        db.add(item)
        db.commit()

        resp = client.post("/product-db/api/quotations", json={
            "solution_id": sol.id,
            "title": "从方案创建",
        })
        assert resp.status_code in (200, 201)
        qt = resp.json()["quotation"]
        assert qt["client_name"] == "客户甲"

    def test_tax_rate_defaults_to_13(self, db):
        """报价单税率统一 13%（导出的 xlsx 备注行会写「税率 N%」，客户看得到）"""
        qt = client.post("/product-db/api/quotations", json={"title": "税率"}).json()["quotation"]
        assert qt["tax_rate"] == 13

    def test_tax_rate_can_be_overridden(self, db):
        """显式传入时仍以传入值为准（不要把 13 写死）"""
        qt = client.post("/product-db/api/quotations",
                         json={"title": "税率", "tax_rate": 6}).json()["quotation"]
        assert qt["tax_rate"] == 6

    def test_export_note_shows_tax_rate(self, db):
        """导出 xlsx 的备注行要写明含税口径与税率（口径：单价即含税价，不再二次计税）"""
        import io
        import openpyxl

        qt = client.post("/product-db/api/quotations", json={"title": "导出税率"}).json()["quotation"]
        res = client.get(f"/product-db/api/quotations/{qt['id']}/export-xlsx")
        assert res.status_code == 200
        ws = openpyxl.load_workbook(io.BytesIO(res.content)).active
        texts = [str(c.value) for row in ws.iter_rows() for c in row if c.value]
        notes = [t for t in texts if t.startswith("注：")]
        assert notes, texts[-4:]
        assert "含税价" in notes[0] and "13% 增值税" in notes[0], notes[0]

    def test_update_quotation(self, db):
        resp = client.post("/product-db/api/quotations", json={"title": "原始报价"})
        qt_id = resp.json()["quotation"]["id"]

        resp = client.put(f"/product-db/api/quotations/{qt_id}", json={"title": "更新报价"})
        assert resp.status_code in (200, 201)
        assert resp.json()["quotation"]["title"] == "更新报价"

    def test_delete_quotation(self, db):
        resp = client.post("/product-db/api/quotations", json={"title": "删除报价"})
        qt_id = resp.json()["quotation"]["id"]

        resp = client.delete(f"/product-db/api/quotations/{qt_id}")
        assert resp.status_code in (200, 201)

    def test_quotation_items_crud(self, db):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)
        p.base_price = 100
        db.commit()

        resp = client.post("/product-db/api/quotations", json={"title": "报价项测试"})
        qt_id = resp.json()["quotation"]["id"]

        # Add item
        resp = client.post(f"/product-db/api/quotations/{qt_id}/items", json={
            "product_id": p.id, "quantity": 3, "unit_price": 100,
        })
        assert resp.status_code in (200, 201)
        item_id = resp.json()["item"]["id"]

        # List
        resp = client.get(f"/product-db/api/quotations/{qt_id}/items")
        assert resp.status_code == 200
        assert len(resp.json()["items"]) == 1

        # Update
        resp = client.put(f"/product-db/api/quotations/{qt_id}/items/{item_id}",
                          json={"quantity": 5})
        assert resp.status_code == 200

        # Delete
        resp = client.delete(f"/product-db/api/quotations/{qt_id}/items/{item_id}")
        assert resp.status_code == 200

    def test_export_xlsx(self, db):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)
        p.base_price = 200
        db.commit()

        resp = client.post("/product-db/api/quotations", json={"title": "导出测试"})
        qt_id = resp.json()["quotation"]["id"]
        client.post(f"/product-db/api/quotations/{qt_id}/items", json={
            "product_id": p.id, "quantity": 1, "unit_price": 200,
        })

        resp = client.get(f"/product-db/api/quotations/{qt_id}/export-xlsx")
        assert resp.status_code in (200, 201)
        assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")


# ============================================================
# BOM Templates
# ============================================================
class TestBOMTemplates:
    def test_create_and_list(self, db):
        resp = client.post("/product-db/api/bom-templates", json={
            "name": "标准模板", "snapshot": {"cells": {"A1": {"v": "标题"}}},
        })
        assert resp.status_code in (200, 201)
        t_id = resp.json()["template"]["id"]

        resp = client.get("/product-db/api/bom-templates")
        assert resp.status_code == 200
        assert len(resp.json()["templates"]) >= 1

    def test_get_template(self, db):
        resp = client.post("/product-db/api/bom-templates", json={
            "name": "查询模板", "snapshot": {},
        })
        t_id = resp.json()["template"]["id"]

        resp = client.get(f"/product-db/api/bom-templates/{t_id}")
        assert resp.status_code in (200, 201)
        assert resp.json()["template"]["name"] == "查询模板"

    def test_update_template(self, db):
        resp = client.post("/product-db/api/bom-templates", json={
            "name": "旧模板", "snapshot": {},
        })
        t_id = resp.json()["template"]["id"]

        resp = client.put(f"/product-db/api/bom-templates/{t_id}", json={"name": "新模板"})
        assert resp.status_code in (200, 201)
        assert resp.json()["template"]["name"] == "新模板"

    def test_delete_template(self, db):
        resp = client.post("/product-db/api/bom-templates", json={
            "name": "删除模板", "snapshot": {},
        })
        t_id = resp.json()["template"]["id"]

        resp = client.delete(f"/product-db/api/bom-templates/{t_id}")
        assert resp.status_code in (200, 201)

    def test_duplicate_template(self, db):
        resp = client.post("/product-db/api/bom-templates", json={
            "name": "原始模板", "snapshot": {"cells": {"A1": {"v": "x"}}},
        })
        t_id = resp.json()["template"]["id"]

        resp = client.post(f"/product-db/api/bom-templates/{t_id}/duplicate")
        assert resp.status_code in (200, 201)
        assert "副本" in resp.json()["template"]["name"]


# ============================================================
# BOM Snapshots
# ============================================================
class TestBOMSnapshots:
    def test_get_bom_snapshot_generates_from_template(self, db):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)

        # Create default template
        client.post("/product-db/api/bom-templates", json={
            "name": "默认模板", "snapshot": {"cells": {"A1": {"v": "BOM"}}}, "is_default": True,
        })

        # Create solution with item
        resp = client.post("/product-db/api/solutions", json={"name": "BOM方案"})
        sol_id = resp.json()["solution"]["id"]
        client.post(f"/product-db/api/solutions/{sol_id}/items", json={
            "product_id": p.id, "quantity": 2, "unit_price": 100,
        })

        # Get snapshot (auto-generated)
        resp = client.get(f"/product-db/api/solutions/{sol_id}/bom-snapshot")
        assert resp.status_code in (200, 201)
        assert "bom_snapshot" in resp.json()
        # Header row is now overwritten with unified layout
        assert resp.json()["bom_snapshot"]["snapshot"]["cells"]["A1"]["v"] == "#"
        # Data row should have product info
        cells = resp.json()["bom_snapshot"]["snapshot"]["cells"]
        assert cells["B3"]["v"] == "测试产品"  # product name in data row

    def test_save_bom_snapshot(self, db):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)

        resp = client.post("/product-db/api/solutions", json={"name": "快照方案"})
        sol_id = resp.json()["solution"]["id"]

        # Save snapshot
        resp = client.put(f"/product-db/api/solutions/{sol_id}/bom-snapshot", json={
            "snapshot": {"cells": {"A1": {"v": "自定义内容"}}},
        })
        assert resp.status_code in (200, 201)

        # Read back
        resp = client.get(f"/product-db/api/solutions/{sol_id}/bom-snapshot")
        assert resp.status_code == 200
        assert resp.json()["bom_snapshot"]["snapshot"]["cells"]["A1"]["v"] == "自定义内容"

    def test_save_as_template(self, db):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)

        resp = client.post("/product-db/api/solutions", json={"name": "另存模板方案"})
        sol_id = resp.json()["solution"]["id"]

        client.put(f"/product-db/api/solutions/{sol_id}/bom-snapshot", json={
            "snapshot": {"cells": {"A1": {"v": "另存测试"}}},
        })

        resp = client.post(f"/product-db/api/solutions/{sol_id}/bom-snapshot/save-as-template", json={
            "name": "从快照创建模板",
        })
        assert resp.status_code in (200, 201)
        assert resp.json()["template"]["name"] == "从快照创建模板"

    def test_export_bom_xlsx(self, db):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)

        resp = client.post("/product-db/api/solutions", json={"name": "导出BOM方案"})
        sol_id = resp.json()["solution"]["id"]
        client.post(f"/product-db/api/solutions/{sol_id}/items", json={
            "product_id": p.id, "quantity": 1, "unit_price": 500,
        })

        resp = client.get(f"/product-db/api/solutions/{sol_id}/bom-snapshot/export-xlsx")
        assert resp.status_code in (200, 201)
        assert resp.headers["content-type"].startswith("application/vnd.openxmlformats")


# ============================================================
# Dictionaries
# ============================================================
class TestDictionaries:
    def test_comm_methods(self, db):
        db.add(DictCommMethod(name="Ethernet", method_type="wired"))
        db.commit()
        resp = client.get("/product-db/api/dicts/comm-methods")
        assert resp.status_code == 200

    def test_comm_protocols(self, db):
        db.add(DictCommProtocol(name="MQTT"))
        db.commit()
        resp = client.get("/product-db/api/dicts/comm-protocols")
        assert resp.status_code == 200

    def test_manufacturers(self, db):
        _seed_manufacturer(db)
        resp = client.get("/product-db/api/dicts/manufacturers")
        assert resp.status_code == 200


# ============================================================
# AI Chat
# ============================================================
class TestAI:
    def test_chat_no_input(self):
        resp = client.post("/product-db/api/ai/chat", json={"input": ""})
        assert resp.status_code == 400

    def test_conversations(self):
        resp = client.get("/product-db/api/ai/conversations")
        assert resp.status_code in (200, 201)


# ============================================================
# Settings
# ============================================================
class TestSettings:
    def test_list(self):
        resp = client.get("/product-db/api/settings")
        assert resp.status_code == 200


# ============================================================
# Import Preview
# ============================================================
class TestImport:
    def test_import_preview_no_file(self):
        resp = client.post("/product-db/api/products/import-preview")
        assert resp.status_code == 422


# ============================================================
# Spec Validation (unit test)
# ============================================================
class TestSpecValidation:
    def test_validate_enum(self, db):
        cat = _seed_category(db)
        sd = CategorySpecDefinition(
            category_id=cat.id, spec_key="ip_rating", display_name="IP等级",
            spec_type="enum", options=["IP30", "IP65", "IP67"],
        )
        db.add(sd)
        db.commit()

        from app.services.spec_service import validate_specs
        errors = validate_specs({"ip_rating": "IP65"}, [sd])
        assert len(errors) == 0

        errors = validate_specs({"ip_rating": "IP99"}, [sd])
        assert len(errors) == 1

    def test_validate_number(self, db):
        cat = _seed_category(db)
        sd = CategorySpecDefinition(
            category_id=cat.id, spec_key="weight_g", display_name="重量",
            spec_type="number", validation={"min": 0, "max": 50000},
        )
        db.add(sd)
        db.commit()

        from app.services.spec_service import validate_specs
        errors = validate_specs({"weight_g": 500}, [sd])
        assert len(errors) == 0

        errors = validate_specs({"weight_g": 99999}, [sd])
        assert len(errors) == 1

    def test_validate_boolean(self, db):
        cat = _seed_category(db)
        sd = CategorySpecDefinition(
            category_id=cat.id, spec_key="has_display", display_name="显示屏",
            spec_type="boolean",
        )
        db.add(sd)
        db.commit()

        from app.services.spec_service import validate_specs
        errors = validate_specs({"has_display": True}, [sd])
        assert len(errors) == 0

        errors = validate_specs({"has_display": "yes"}, [sd])
        assert len(errors) == 1


class TestFieldVisibility:
    """Test that field visibility settings hide sensitive fields from non-admin users."""

    def _make_user(self, db, username="normal", role="user", can_view_cost=None):
        from app.auth import hash_password
        u = User(username=username, password_hash=hash_password("test123"), role=role,
                 can_view_cost=can_view_cost)
        db.add(u)
        db.commit()
        db.refresh(u)
        return u

    def _auth_for(self, user):
        token = create_token(user.id, user.username)
        return {"Authorization": f"Bearer {token}"}

    def _seed_field_setting(self, db, field_name, user_visible):
        from app.models.field_setting import FieldSetting
        fs = db.query(FieldSetting).filter_by(field_name=field_name).first()
        if fs:
            fs.user_visible = user_visible
        else:
            db.add(FieldSetting(field_name=field_name, user_visible=user_visible))
        db.commit()

    def _clear_field_settings(self, db):
        from app.models.field_setting import FieldSetting
        db.query(FieldSetting).delete()
        db.commit()
        # Clear cache
        from app.services.field_visibility import _cache
        _cache["ts"] = 0
        _cache["data"] = None

    def test_list_hides_cost_price_for_user(self, db):
        """GET /products hides cost_price when field visible=false for non-admin."""
        self._clear_field_settings(db)
        self._seed_field_setting(db, "cost_price", False)
        cat = _seed_category(db)
        _seed_product(db, category_id=cat.id, cost_price=999.99)

        # Admin sees cost_price
        admin = db.query(User).filter_by(username="admin").first()
        res = client.get("/product-db/api/products", headers=self._auth_for(admin))
        assert res.status_code == 200
        assert res.json()["products"][0]["cost_price"] == 999.99

        # Normal user should NOT see cost_price
        user = self._make_user(db)
        res = client.get("/product-db/api/products", headers=self._auth_for(user))
        assert res.status_code == 200
        assert res.json()["products"][0]["cost_price"] is None

    def test_detail_hides_cost_price_for_user(self, db):
        """GET /products/{id} hides cost_price when field visible=false for non-admin."""
        self._clear_field_settings(db)
        self._seed_field_setting(db, "cost_price", False)
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, cost_price=888.88)

        # Admin sees cost_price
        admin = db.query(User).filter_by(username="admin").first()
        res = client.get(f"/product-db/api/products/{p.id}", headers=self._auth_for(admin))
        assert res.status_code == 200
        assert res.json()["product"]["cost_price"] == 888.88

        # Normal user should NOT see cost_price
        user = self._make_user(db)
        res = client.get(f"/product-db/api/products/{p.id}", headers=self._auth_for(user))
        assert res.status_code == 200
        assert res.json()["product"]["cost_price"] is None

    def test_detail_shows_cost_price_when_visible(self, db):
        """GET /products/{id} shows cost_price when field visible=true for non-admin."""
        self._clear_field_settings(db)
        self._seed_field_setting(db, "cost_price", True)
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, cost_price=777.77)

        user = self._make_user(db)
        res = client.get(f"/product-db/api/products/{p.id}", headers=self._auth_for(user))
        assert res.status_code == 200
        assert res.json()["product"]["cost_price"] == 777.77

    def test_quotation_snapshot_hides_cost_price_for_user(self, db):
        """Quotation item snapshot + BOM must hide cost_price when field visible=false."""
        self._clear_field_settings(db)
        self._seed_field_setting(db, "cost_price", False)
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, base_price=100, cost_price=66.6)
        admin = db.query(User).filter_by(username="admin").first()
        user = self._make_user(db)

        # User creates own quotation with a product item (snapshot contains cost internally)
        res = client.post("/product-db/api/quotations", json={"title": "T"}, headers=self._auth_for(user))
        assert res.status_code == 201
        qid = res.json()["quotation"]["id"]
        res = client.post(
            f"/product-db/api/quotations/{qid}/items",
            json={"product_id": p.id, "quantity": 1, "unit_price": 50},
            headers=self._auth_for(user),
        )
        assert res.status_code == 201

        # User must NOT see cost_price in the quotation snapshot
        res = client.get(f"/product-db/api/quotations/{qid}", headers=self._auth_for(user))
        assert res.status_code == 200
        snap = res.json()["quotation"]["items"][0]["product_snapshot"]
        assert "cost_price" not in snap

        # Items endpoint must not leak cost_price either
        res = client.get(f"/product-db/api/quotations/{qid}/items", headers=self._auth_for(user))
        assert res.status_code == 200
        snap = res.json()["items"][0]["product_snapshot"]
        assert "cost_price" not in snap

        # BOM rows must not expose cost
        res = client.get(f"/product-db/api/quotations/{qid}/bom", headers=self._auth_for(user))
        assert res.status_code == 200
        assert res.json()["rows"][0].get("cost") is None

        # Admin still sees cost_price
        res = client.get(f"/product-db/api/quotations/{qid}", headers=self._auth_for(admin))
        assert res.status_code == 200
        snap = res.json()["quotation"]["items"][0]["product_snapshot"]
        assert snap.get("cost_price") == 66.6

    def test_bom_snapshot_and_export_hide_cost_for_user(self, db):
        """BOM 快照与 xlsx 导出也必须遵守字段可见性。

        回归：`get_bom_snapshot` 原样返回快照、`export_bom_xlsx` 在「有快照」的分支
        里完全不看 show_cost，而快照的 J 列固定写死每个产品的 cost_price →
        管理员关掉成本可见性后，普通用户照样能拿到逐条成本。
        """
        import io
        import openpyxl

        self._clear_field_settings(db)
        self._seed_field_setting(db, "cost_price", False)
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, base_price=100, cost_price=66.6)
        user = self._make_user(db)
        admin = db.query(User).filter_by(username="admin").first()

        res = client.post("/product-db/api/solutions", json={"name": "BOM 成本"},
                          headers=self._auth_for(user))
        assert res.status_code in (200, 201), res.text
        sid = res.json()["solution"]["id"]
        res = client.post(f"/product-db/api/solutions/{sid}/items",
                          json={"product_id": p.id, "quantity": 2, "unit_price": 50},
                          headers=self._auth_for(user))
        assert res.status_code in (200, 201), res.text

        # 快照（首次读会生成并落库）：普通用户看不到 J 列，其它列不受影响
        res = client.get(f"/product-db/api/solutions/{sid}/bom-snapshot",
                         headers=self._auth_for(user))
        assert res.status_code == 200
        cells = res.json()["bom_snapshot"]["snapshot"]["cells"]
        assert not [k for k in cells if k.startswith("J")], "快照不得含成本列"
        assert cells.get("B1", {}).get("v") == "产品名称", "其它列必须保留"

        # 导出走的是「已有快照」分支，正是历史上漏过滤的那条
        res = client.get(f"/product-db/api/solutions/{sid}/bom-snapshot/export-xlsx",
                         headers=self._auth_for(user))
        assert res.status_code == 200
        ws = openpyxl.load_workbook(io.BytesIO(res.content)).active
        assert ws["J3"].value in (None, ""), "导出不得含成本值"
        assert ws["B3"].value == p.name, "导出仍需包含产品行"

        # admin 照旧能拿到成本
        res = client.get(f"/product-db/api/solutions/{sid}/bom-snapshot",
                         headers=self._auth_for(admin))
        assert res.status_code == 200
        admin_cells = res.json()["bom_snapshot"]["snapshot"]["cells"]
        assert any(k.startswith("J") for k in admin_cells), "admin 仍应看到成本列"
        assert admin_cells["J3"]["v"] == 66.6

    def test_admin_always_sees_cost_price(self, db):
        """Admin sees cost_price regardless of field visibility setting."""
        self._clear_field_settings(db)
        self._seed_field_setting(db, "cost_price", False)
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, cost_price=555.55)

        admin = db.query(User).filter_by(username="admin").first()
        res = client.get(f"/product-db/api/products/{p.id}", headers=self._auth_for(admin))
        assert res.status_code == 200
        assert res.json()["product"]["cost_price"] == 555.55

    def test_session_returns_field_visibility_for_user(self, db):
        """GET /auth/session returns field_visibility for non-admin, empty for admin."""
        self._clear_field_settings(db)
        self._seed_field_setting(db, "cost_price", False)
        self._seed_field_setting(db, "manufacturer_name", True)

        # Non-admin gets field_visibility dict
        user = self._make_user(db)
        res = client.get("/product-db/api/auth/session", headers=self._auth_for(user))
        assert res.status_code == 200
        fv = res.json()["field_visibility"]
        assert fv["cost_price"] is False
        assert fv["manufacturer_name"] is True

        # Admin gets empty dict (sees everything)
        admin = db.query(User).filter_by(username="admin").first()
        res = client.get("/product-db/api/auth/session", headers=self._auth_for(admin))
        assert res.status_code == 200
        assert res.json()["field_visibility"] == {}

    def test_no_field_settings_defaults_visible(self, db):
        """When no FieldSetting rows exist, all fields default to visible."""
        self._clear_field_settings(db)
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, cost_price=333.33)

        # Admin endpoint auto-creates defaults, but regular endpoints
        # just pass through (apply_field_visibility has nothing to hide)
        user = self._make_user(db)
        res = client.get(f"/product-db/api/products/{p.id}", headers=self._auth_for(user))
        assert res.status_code == 200
        # Without field settings, no fields are filtered
        assert res.json()["product"]["cost_price"] == 333.33

    # ── 成本可见性统一判定（R36）：方案路径此前完全没接字段可见性 ──
    #
    # 背景：`cost_price` / `product_cost_price` / `total_cost` 是三个不同的键名，
    # 而字段可见性只认第一个。方案接口直接返回 `Solution.to_dict()`（含方案总成本）
    # 与 `SolutionItem.to_dict()`（含每项成本），普通用户拿到方案即拿到全部成本。

    def _seed_solution_with_cost(self, db, user, cost=66.6, qty=2):
        # 复用已存在的品类，避免同一个测试里建两次触发 slug 唯一约束
        cat = db.query(Category).filter_by(slug="test-cat").first() or _seed_category(db)
        p = _seed_product(db, category_id=cat.id, base_price=100, cost_price=cost)
        res = client.post("/product-db/api/solutions", json={"name": "成本方案"},
                          headers=self._auth_for(user))
        assert res.status_code in (200, 201), res.text
        sid = res.json()["solution"]["id"]
        res = client.post(f"/product-db/api/solutions/{sid}/items",
                          json={"product_id": p.id, "quantity": qty, "unit_price": 50},
                          headers=self._auth_for(user))
        assert res.status_code in (200, 201), res.text
        return sid, res.json()["item"]

    def test_solution_cost_hidden_from_user(self, db):
        """全局关闭成本可见性时，方案的 6 个返回点都不得泄漏成本。"""
        self._clear_field_settings(db)
        self._seed_field_setting(db, "cost_price", False)
        user = self._make_user(db)
        sid, item = self._seed_solution_with_cost(db, user)

        # 1) 新增条目接口（历史泄漏点）
        assert "product_cost_price" not in item, "新增条目返回不得含成本"

        # 2) 方案详情
        res = client.get(f"/product-db/api/solutions/{sid}", headers=self._auth_for(user))
        assert res.status_code == 200
        sol = res.json()["solution"]
        assert "total_cost" not in sol, "方案总成本不得外泄"
        assert all("product_cost_price" not in it for it in sol["items"])

        # 3) 方案列表（同样带 items）
        res = client.get("/product-db/api/solutions", headers=self._auth_for(user))
        assert res.status_code == 200
        for s in res.json()["solutions"]:
            assert "total_cost" not in s
            for it in s["items"]:
                assert "product_cost_price" not in it

        # 4) 条目列表
        res = client.get(f"/product-db/api/solutions/{sid}/items", headers=self._auth_for(user))
        assert res.status_code == 200
        assert all("product_cost_price" not in it for it in res.json()["items"])

        # 5) 更新方案 / 6) 更新条目
        res = client.put(f"/product-db/api/solutions/{sid}", json={"name": "改名"},
                         headers=self._auth_for(user))
        assert res.status_code == 200
        assert "total_cost" not in res.json()["solution"]

        item_id = client.get(f"/product-db/api/solutions/{sid}/items",
                             headers=self._auth_for(user)).json()["items"][0]["id"]
        res = client.put(f"/product-db/api/solutions/{sid}/items/{item_id}",
                         json={"quantity": 3}, headers=self._auth_for(user))
        assert res.status_code == 200
        assert "product_cost_price" not in res.json()["item"]

    def test_solution_cost_visible_when_global_on_and_for_admin(self, db):
        """全局开关打开时普通用户可见；全局关掉后 admin 仍可见。"""
        self._clear_field_settings(db)
        self._seed_field_setting(db, "cost_price", True)
        user = self._make_user(db)
        sid, item = self._seed_solution_with_cost(db, user)

        assert item["product_cost_price"] == 66.6
        res = client.get(f"/product-db/api/solutions/{sid}", headers=self._auth_for(user))
        assert res.json()["solution"]["total_cost"] == 133.2

        # 把全局开关关掉（_clear_field_settings 会一并清 30s 缓存）
        self._clear_field_settings(db)
        self._seed_field_setting(db, "cost_price", False)

        admin = db.query(User).filter_by(username="admin").first()
        res = client.get(f"/product-db/api/solutions/{sid}", headers=self._auth_for(admin))
        assert res.status_code == 200
        assert res.json()["solution"]["total_cost"] == 133.2
        assert res.json()["solution"]["items"][0]["product_cost_price"] == 66.6

    def test_bom_snapshot_save_response_hides_cost(self, db):
        """PUT /bom-snapshot 的响应也不得回传成本列（此前直接返回落库后的快照）。"""
        self._clear_field_settings(db)
        self._seed_field_setting(db, "cost_price", False)
        user = self._make_user(db)
        admin = db.query(User).filter_by(username="admin").first()
        sid, _item = self._seed_solution_with_cost(db, user)

        # 由 admin 触发生成快照：落库的快照里带 J 列成本
        res = client.get(f"/product-db/api/solutions/{sid}/bom-snapshot",
                         headers=self._auth_for(admin))
        assert res.status_code == 200
        assert any(k.startswith("J") for k in res.json()["bom_snapshot"]["snapshot"]["cells"])

        # 非管理员保存时传空快照 → 服务端保留原快照（含 J），响应必须裁掉成本
        res = client.put(f"/product-db/api/solutions/{sid}/bom-snapshot",
                         json={"snapshot": {}}, headers=self._auth_for(user))
        assert res.status_code == 200
        cells = res.json()["bom_snapshot"]["snapshot"]["cells"]
        assert not [k for k in cells if k.startswith("J")], "保存响应不得回传成本列"
        assert cells.get("B1", {}).get("v") == "产品名称", "其它列必须保留"

    def test_bom_save_keeps_server_cost_for_user(self, db):
        """回归：非管理员保存 BOM 不得把库里的成本列抹成 0。

        前端保存时会无条件回传 J 列，而它来自被裁掉成本列的读取结果（值为 0），
        保存又是整份快照替换 —— 历史上一次保存就把成本清零，之后管理员看到全 0。
        """
        self._clear_field_settings(db)
        self._seed_field_setting(db, "cost_price", False)
        user = self._make_user(db)
        admin = db.query(User).filter_by(username="admin").first()
        sid, _item = self._seed_solution_with_cost(db, user)

        # admin 触发生成快照（落库 J3 = 66.6）
        res = client.get(f"/product-db/api/solutions/{sid}/bom-snapshot",
                         headers=self._auth_for(admin))
        assert res.status_code == 200
        assert res.json()["bom_snapshot"]["snapshot"]["cells"]["J3"]["v"] == 66.6

        # 非管理员按前端行为提交：J3 是 0，同时改了其它列
        payload = {"cells": {
            "B3": {"v": "改名后的产品"}, "E3": {"v": 3}, "F3": {"v": 50},
            "G3": {"v": 100}, "I3": {"v": "改备注"}, "J3": {"v": 0},
        }, "sheet_name": "BOM"}
        res = client.put(f"/product-db/api/solutions/{sid}/bom-snapshot",
                         json={"snapshot": payload}, headers=self._auth_for(user))
        assert res.status_code == 200

        # 库里成本必须还是服务端原值，其它列按客户端提交保存
        cells = client.get(f"/product-db/api/solutions/{sid}/bom-snapshot",
                           headers=self._auth_for(admin)).json()["bom_snapshot"]["snapshot"]["cells"]
        assert cells["J3"]["v"] == 66.6, "保存不得抹掉成本列"
        assert cells["B3"]["v"] == "改名后的产品"
        assert cells["I3"]["v"] == "改备注"

    def test_quotation_bom_save_keeps_server_cost_for_user(self, db):
        """回归：非管理员保存报价单 BOM 不得把快照里的 cost_price 归零。"""
        self._clear_field_settings(db)
        self._seed_field_setting(db, "cost_price", False)
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, base_price=100, cost_price=66.6)
        user = self._make_user(db)
        admin = db.query(User).filter_by(username="admin").first()

        qid = client.post("/product-db/api/quotations", json={"title": "T"},
                          headers=self._auth_for(user)).json()["quotation"]["id"]
        assert client.post(f"/product-db/api/quotations/{qid}/items",
                           json={"product_id": p.id, "quantity": 1, "unit_price": 50},
                           headers=self._auth_for(user)).status_code == 201

        # 前端保存 BOM 的载荷：cost 来自读取结果（非管理员拿到的是 None）
        res = client.put(f"/product-db/api/quotations/{qid}/bom",
                         json={"rows": [{"name": p.name, "sku": p.sku or "SKU-1", "model": "",
                                         "description": "", "qty": 2, "price": 50,
                                         "discount": 100, "remark": "", "cost": None}]},
                         headers=self._auth_for(user))
        assert res.status_code == 200

        snap = client.get(f"/product-db/api/quotations/{qid}",
                          headers=self._auth_for(admin)).json()["quotation"]["items"][0]["product_snapshot"]
        assert snap.get("cost_price") == 66.6, "保存不得把成本归零"
        assert snap.get("name") == p.name, "其它字段仍按客户端提交保存"

    # ── 按用户的成本可见性覆盖（users.can_view_cost，三态）──

    def test_user_override_allows_cost_when_global_off(self, db):
        """全局关掉成本，仍可给单个用户单独开放（覆盖优先于全局）。"""
        self._clear_field_settings(db)
        self._seed_field_setting(db, "cost_price", False)
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, cost_price=4321.0)
        user = self._make_user(db, username="vip", can_view_cost=True)

        res = client.get(f"/product-db/api/products/{p.id}", headers=self._auth_for(user))
        assert res.status_code == 200
        assert res.json()["product"]["cost_price"] == 4321.0

        sid, item = self._seed_solution_with_cost(db, user, cost=4321.0)
        assert item["product_cost_price"] == 4321.0
        assert client.get(f"/product-db/api/solutions/{sid}",
                          headers=self._auth_for(user)).json()["solution"]["total_cost"] == 8642.0

    def test_user_override_denies_cost_when_global_on(self, db):
        """反向覆盖：全局开着，也可以对某个人单独禁止。"""
        self._clear_field_settings(db)
        self._seed_field_setting(db, "cost_price", True)
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, cost_price=1234.0)
        user = self._make_user(db, username="blocked", can_view_cost=False)

        res = client.get(f"/product-db/api/products/{p.id}", headers=self._auth_for(user))
        assert res.status_code == 200
        assert res.json()["product"]["cost_price"] is None

        sid, item = self._seed_solution_with_cost(db, user, cost=1234.0)
        assert "product_cost_price" not in item
        sol = client.get(f"/product-db/api/solutions/{sid}",
                         headers=self._auth_for(user)).json()["solution"]
        assert "total_cost" not in sol

    def test_session_reports_effective_can_view_cost(self, db):
        """GET /auth/session 暴露生效后的成本可见性（admin / 覆盖 / 全局三者合一）。"""
        self._clear_field_settings(db)
        self._seed_field_setting(db, "cost_price", False)

        def session_flag(u):
            res = client.get("/product-db/api/auth/session", headers=self._auth_for(u))
            assert res.status_code == 200
            return res.json()["can_view_cost"]

        admin = db.query(User).filter_by(username="admin").first()
        assert session_flag(admin) is True, "admin 恒可见"
        assert session_flag(self._make_user(db, username="follow")) is False, "NULL 跟随全局"
        assert session_flag(self._make_user(db, username="allow", can_view_cost=True)) is True
        assert session_flag(self._make_user(db, username="deny", can_view_cost=False)) is False

    def test_admin_sets_can_view_cost_three_state(self, db):
        """管理端可三态设置：允许 / 禁止 / 改回跟随全局（显式 null）。"""
        admin = db.query(User).filter_by(username="admin").first()
        target = self._make_user(db, username="target")
        url = f"/product-db/api/admin/users/{target.id}"

        for value in (True, False):
            res = client.put(url, json={"can_view_cost": value}, headers=self._auth_for(admin))
            assert res.status_code == 200
            assert res.json()["user"]["can_view_cost"] is value

        # 改回跟随全局必须真的生效 —— apply_partial_update 会跳过 None，
        # 若把 can_view_cost 交给它，这里会一直是 False（回归点）
        res = client.put(url, json={"can_view_cost": None}, headers=self._auth_for(admin))
        assert res.status_code == 200
        assert res.json()["user"]["can_view_cost"] is None

        # 列表接口也要带上该字段，管理页才能显示当前状态
        users = client.get("/product-db/api/admin/users", headers=self._auth_for(admin)).json()["users"]
        assert next(u for u in users if u["id"] == target.id)["can_view_cost"] is None

        # 不带该字段的普通更新不得把它重置
        assert client.put(url, json={"can_view_cost": False},
                          headers=self._auth_for(admin)).status_code == 200
        res = client.put(url, json={"email": "x@example.com"}, headers=self._auth_for(admin))
        assert res.status_code == 200
        assert res.json()["user"]["can_view_cost"] is False, "未提及该字段时不应被重置"

    def test_product_export_hides_cost_column(self, db):
        """产品导出 xlsx 里，普通用户不该看到成本 —— 值不写、「成本」表头也不写（R46）。"""
        import io
        import openpyxl

        self._clear_field_settings(db)
        self._seed_field_setting(db, "cost_price", False)
        cat = _seed_category(db)
        _seed_product(db, category_id=cat.id, cost_price=66.6)
        user = self._make_user(db)
        admin = db.query(User).filter_by(username="admin").first()

        res = client.get("/product-db/api/products/export", headers=self._auth_for(user))
        assert res.status_code == 200
        ws = openpyxl.load_workbook(io.BytesIO(res.content)).active
        # 第 3 行是表头，数据从第 4 行开始（enumerate(..., 1) → 3 + idx）
        assert ws.cell(row=3, column=13).value is None, "「成本」表头也不该出现"
        assert ws.cell(row=4, column=13).value is None, "导出不得含成本值"
        assert ws.cell(row=4, column=2).value, "导出仍需包含产品行"

        res = client.get("/product-db/api/products/export", headers=self._auth_for(admin))
        assert res.status_code == 200
        ws = openpyxl.load_workbook(io.BytesIO(res.content)).active
        assert ws.cell(row=4, column=13).value == 66.6, "admin 仍应看到成本列"


class TestPerimeterHardening:
    """收敛几处「对外暴露面过大 / 静默失败」的审查项。"""

    def _make_user(self, db, username, role="user"):
        u = User(username=username, password_hash=hash_password("test123"), role=role)
        db.add(u)
        db.commit()
        db.refresh(u)
        return u

    def _auth_for(self, user):
        return {"Authorization": f"Bearer {create_token(user.id, user.username)}"}

    def test_ai_stats_hides_global_totals_for_non_admin(self, db):
        """回归：/ai/stats 曾把全站 AI 用量（total / total_tokens_*）返回给任意登录用户。"""
        admin = db.query(User).filter_by(username="admin").first()
        res = client.get("/product-db/api/ai/stats", headers=self._auth_for(admin))
        assert res.status_code == 200
        assert "total" in res.json(), "admin 应拿到全局口径"

        user = self._make_user(db, "stats_user")
        res = client.get("/product-db/api/ai/stats", headers=self._auth_for(user))
        assert res.status_code == 200
        body = res.json()
        assert "total" not in body and "total_tokens_in" not in body, "非 admin 不得看到全局口径"
        # 侧边栏用的用户级字段必须保留，否则界面会空
        assert "user_count" in body and "user_tokens_in" in body

    def test_import_preview_rejects_oversize_upload(self, db, monkeypatch):
        """回归：import-preview 曾把整个文件读进内存且无上限（普通用户可打满内存）。"""
        import io
        from app.config import settings

        user = self._make_user(db, "import_user")
        monkeypatch.setattr(settings, "FILE_MAX_SIZE", 1024, raising=False)

        res = client.post(
            "/product-db/api/products/import-preview",
            files={"file": ("big.xlsx", io.BytesIO(b"x" * 4096),
                            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=self._auth_for(user),
        )
        assert res.status_code == 400
        assert "too large" in res.json()["detail"].lower()

    def test_product_view_count_is_incremented_atomically(self, db):
        """浏览量改为 SQL 层自增后仍要正确累加（view_count 为 NULL 时不能变 NULL）。"""
        admin = db.query(User).filter_by(username="admin").first()
        cat = _seed_category(db)
        p = _seed_product(db, name="浏览量产品", model="VC-1", category_id=cat.id)
        p.view_count = None
        db.commit()

        for expected in (1, 2):
            res = client.get(f"/product-db/api/products/{p.id}", headers=self._auth_for(admin))
            assert res.status_code == 200
            assert res.json()["product"]["view_count"] == expected


class TestRateLimitExemption:
    """探针与静态资源不得被全局限流计数（R38）。

    背景：slowapi 的中间件按「解析到的 handler 名字」匹配限流规则，而 SPA catch-all
    `/product-db/{full_path:path}` 注册在最后、遮蔽了所有 handler → 函数级
    `@limiter.exempt` / `@limiter.limit` 全部无效，只剩 default_limits 生效。
    后果（2026-09 生产实测）：探针 2 请求/2 分钟 = 1440 次/天，3.3 小时打满
    200/天配额，此后全天 429 —— 09-18~09-20 每天 380+ 次误报「服务不可用」，
    还会污染探针状态机（真宕机时不再产生新告警）。
    """

    def test_exempt_paths(self):
        from app.main import rate_limit_exempt
        for p in ("/product-db/api/health", "/product-db", "/product-db/",
                  "/product-db/assets/main.js", "/product-db/api/uploads",
                  "/product-db/api/uploads/a.png", "/api/uploads/a.png"):
            assert rate_limit_exempt(p), f"{p} 应豁免限流"

    def test_api_paths_still_limited(self):
        from app.main import rate_limit_exempt
        for p in ("/product-db/api/products", "/product-db/api/auth/login",
                  "/product-db/api/solutions/1/items", "/product-db/login",
                  "/product-db/api/ai/chat"):
            assert not rate_limit_exempt(p), f"{p} 不应豁免"

    def test_health_survives_more_than_daily_quota(self):
        """连打 250 次（超过旧默认 200/天）必须全部 200 —— 旧实现第 201 次起 429。"""
        codes = {client.get("/product-db/api/health").status_code for _ in range(250)}
        assert codes == {200}, f"健康接口被限流了: {sorted(codes)}"


class TestImportConfirm:
    """补齐此前**零覆盖**的写路径：POST /products/import-confirm（批量导入落库）。

    审查发现只有 import-preview 的拒绝路径被覆盖，真正写库的这条从未验证过。
    顺带修掉：products.category_id 是 NOT NULL，品类名匹配不到时插 None 会让整批
    在 commit 处 IntegrityError → 500（且无任何可操作提示、整批数据丢失）。
    """

    def _auth_for(self, user):
        return {"Authorization": f"Bearer {create_token(user.id, user.username)}"}

    def test_import_confirm_creates_products_and_maps_fields(self, db):
        admin = db.query(User).filter_by(username="admin").first()
        headers = self._auth_for(admin)
        cat = _seed_category(db, name="导入品类", slug="import-cat")
        cat2 = _seed_category(db, name="无线网关", slug="import-cat2")

        res = client.post("/product-db/api/products/import-confirm", json={
            "mapping": {"0": "name", "1": "model", "2": "price", "3": "category"},
            "rows": [
                ["导入产品A", "IA-1", "12.5", "导入品类"],
                ["导入产品B", "IB-1", "8", "无线"],      # 走模糊匹配分支
            ],
        }, headers=headers)
        assert res.status_code == 200, res.text
        assert res.json()["imported"] == 2

        p1 = db.query(Product).filter_by(model="IA-1").first()
        assert p1 is not None and p1.name == "导入产品A"
        assert float(p1.base_price) == 12.5
        assert p1.category_id == cat.id, "品类名应映射到对应 category_id"

        p2 = db.query(Product).filter_by(model="IB-1").first()
        assert p2 is not None and p2.category_id == cat2.id, "应支持品类名模糊匹配"

    def test_import_confirm_rejects_unmatched_category_without_writing(self, db):
        """回归：匹配不到的品类名曾导致整批 IntegrityError → 500（生产实测）。"""
        admin = db.query(User).filter_by(username="admin").first()
        before = db.query(Product).count()

        res = client.post("/product-db/api/products/import-confirm", json={
            "mapping": {"0": "name", "1": "category"},
            "rows": [["回归产品X", "不存在的品类XYZ"]],
        }, headers=self._auth_for(admin))

        assert res.status_code == 400, f"应给出可操作的 400，而不是 500：{res.status_code}"
        assert "不存在的品类XYZ" in res.json()["detail"]
        assert db.query(Product).count() == before, "校验失败时不得写入任何产品"

    def test_import_confirm_requires_category_column(self, db):
        """Excel 没有品类列时也必须明确报错，而不是 500。"""
        admin = db.query(User).filter_by(username="admin").first()
        res = client.post("/product-db/api/products/import-confirm", json={
            "mapping": {"0": "name", "1": "model"},
            "rows": [["只有名称的产品", "NC-1"]],
        }, headers=self._auth_for(admin))
        assert res.status_code == 400
        assert "品类" in res.json()["detail"]

    def test_import_confirm_skips_empty_rows(self, db):
        """空行按既有语义跳过；全部为空则明确报错而不是返回 imported=0。"""
        admin = db.query(User).filter_by(username="admin").first()
        headers = self._auth_for(admin)
        _seed_category(db, name="跳过测试品类", slug="skip-cat")

        res = client.post("/product-db/api/products/import-confirm", json={
            "mapping": {"0": "name", "1": "category"},
            "rows": [["有效产品", "跳过测试品类"], ["", ""]],
        }, headers=headers)
        assert res.status_code == 200
        assert res.json()["imported"] == 1, "空行应被跳过、其余行正常导入"

        res = client.post("/product-db/api/products/import-confirm", json={
            "mapping": {"0": "name", "1": "category"},
            "rows": [["", ""]],
        }, headers=headers)
        assert res.status_code == 400

    def test_import_confirm_rejects_empty_payload(self, db):
        admin = db.query(User).filter_by(username="admin").first()
        res = client.post("/product-db/api/products/import-confirm",
                          json={"mapping": {}, "rows": []}, headers=self._auth_for(admin))
        assert res.status_code == 400


class TestOwnershipGaps:
    """Ownership checks that were missing: dependencies, export, compare."""

    def _make_user(self, db, username):
        u = User(username=username, password_hash=hash_password("test123"), role="user")
        db.add(u)
        db.commit()
        db.refresh(u)
        return u

    def _auth_for(self, user):
        return {"Authorization": f"Bearer {create_token(user.id, user.username)}"}

    def test_non_admin_cannot_modify_others_product_dependencies(self, db):
        """Dependency create/update/delete must enforce product ownership."""
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)  # created_by NULL → admin/legacy
        # Give the product an explicit owner for a clean 403 scenario
        p.created_by = 1
        db.commit()

        user = self._make_user(db, "dep_user")
        headers = self._auth_for(user)

        res = client.post(
            f"/product-db/api/products/{p.id}/dependencies",
            json={"depends_on_category_id": cat.id, "dependency_type": "required"},
            headers=headers,
        )
        assert res.status_code == 403, res.text

    def test_non_admin_cannot_update_or_delete_others_dependencies(self, db):
        """Existing dependency owned by another user's product must be 403."""
        from app.models.dependency import ProductDependency
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)
        p.created_by = 1
        db.commit()
        dep = ProductDependency(product_id=p.id, depends_on_category_id=cat.id, dependency_type="required")
        db.add(dep)
        db.commit()
        db.refresh(dep)

        user = self._make_user(db, "dep_user2")
        headers = self._auth_for(user)

        res = client.put(
            f"/product-db/api/products/{p.id}/dependencies/{dep.id}",
            json={"description": "hacked"},
            headers=headers,
        )
        assert res.status_code == 403, res.text

        res = client.delete(
            f"/product-db/api/products/{p.id}/dependencies/{dep.id}",
            headers=headers,
        )
        assert res.status_code == 403, res.text

    def test_export_respects_ownership_filter(self, db):
        """Export must not include products owned by other regular users."""
        cat = _seed_category(db)
        admin = db.query(User).filter_by(username="admin").first()
        user1 = self._make_user(db, "exp_user1")
        user2 = self._make_user(db, "exp_user2")

        p_own = _seed_product(db, name="user1产品", category_id=cat.id)
        p_own.created_by = user1.id
        db.commit()

        # user2's list must exclude user1's product
        res = client.get("/product-db/api/products?per_page=100", headers=self._auth_for(user2))
        assert res.status_code == 200
        ids = {pp["id"] for pp in res.json()["products"]}
        assert p_own.id not in ids

        # export must exclude it as well
        res = client.get("/product-db/api/products/export", headers=self._auth_for(user2))
        assert res.status_code == 200
        import io, openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        ws = wb.active
        names = []
        for row in ws.iter_rows(min_row=4, values_only=True):
            if row and row[1]:
                names.append(str(row[1]))
        assert "user1产品" not in names, names[:5]

    def test_compare_respects_ownership_filter(self, db):
        """Compare must not return products owned by other regular users."""
        cat = _seed_category(db)
        user1 = self._make_user(db, "cmp_user1")
        user2 = self._make_user(db, "cmp_user2")

        p_a = _seed_product(db, name="A产品", category_id=cat.id)
        p_a.created_by = user1.id
        p_b = _seed_product(db, name="B产品", category_id=cat.id)
        p_b.created_by = user2.id
        db.commit()

        # user2 compares A (not visible) + B (visible) → only 1 visible → 404
        res = client.get(
            f"/product-db/api/products/compare?product_ids={p_a.id},{p_b.id}",
            headers=self._auth_for(user2),
        )
        assert res.status_code == 404, res.text


class TestMasterDataAdminOnly:
    """回归：主数据（品类/厂商/供应商/字典/BOM 模板）的写操作原用
    check_ownership(strict=False)，而它对 created_by IS NULL 的历史行直接放行；
    生产上这些表几乎全是 NULL（manufacturers 48/48、suppliers 55/55、
    dict_sensor_metrics 36/36…）→ 任何登录用户都能改删全站共用数据。
    现改为仅 admin 可写；业务数据（产品/方案/报价单）仍按归属校验。"""

    def _make_user(self, db, username, role="user"):
        u = User(username=username, password_hash=hash_password("test123"), role=role)
        db.add(u)
        db.commit()
        db.refresh(u)
        return u

    def _auth_for(self, user):
        return {"Authorization": f"Bearer {create_token(user.id, user.username)}"}

    def test_non_admin_cannot_write_master_data(self, db):
        cat = _seed_category(db)
        # 模拟生产现状：历史行 created_by 为 NULL
        cat.created_by = None
        db.commit()
        user = self._make_user(db, "md_user")
        headers = self._auth_for(user)

        cases = [
            ("post", "/product-db/api/categories", {"name": "越权品类"}),
            ("put", f"/product-db/api/categories/{cat.id}", {"name": "改名"}),
            ("delete", f"/product-db/api/categories/{cat.id}", None),
            ("post", "/product-db/api/suppliers", {"name": "越权供应商"}),
            ("post", "/product-db/api/dicts/manufacturers", {"name": "越权厂商"}),
            ("delete", "/product-db/api/dicts/sensor-metrics/1", None),
            ("post", "/product-db/api/bom-templates", {"name": "越权模板"}),
        ]
        for method, url, body in cases:
            call = getattr(client, method)
            res = call(url, json=body, headers=headers) if body is not None else call(url, headers=headers)
            assert res.status_code == 403, f"{method.upper()} {url} 应 403，实际 {res.status_code}"

    def test_admin_can_still_write_master_data(self, db):
        admin = db.query(User).filter_by(username="admin").first()
        headers = self._auth_for(admin)
        res = client.post("/product-db/api/categories", json={"name": "管理员建的品类"}, headers=headers)
        assert res.status_code == 201, res.text

    def test_non_admin_business_data_writes_unaffected(self, db):
        """门禁只针对主数据：普通用户仍可创建/修改自己的产品与方案。"""
        cat = _seed_category(db)
        user = self._make_user(db, "biz_user")
        headers = self._auth_for(user)

        res = client.post("/product-db/api/products",
                          json={"name": "普通用户产品", "model": "U-1", "category_id": cat.id},
                          headers=headers)
        assert res.status_code in (200, 201), res.text
        pid = res.json()["product"]["id"]

        res = client.put(f"/product-db/api/products/{pid}", json={"name": "改名后"}, headers=headers)
        assert res.status_code == 200, res.text

        res = client.post("/product-db/api/solutions", json={"name": "普通用户方案"}, headers=headers)
        assert res.status_code in (200, 201), res.text

        res = client.post("/product-db/api/quotations", json={"title": "普通用户报价"}, headers=headers)
        assert res.status_code == 201, res.text


class TestExportFilenames:
    """导出文件名要带客户与项目名（便于区分），且中文名必须正确编码。

    Content-Disposition 双写：`filename=<ASCII 回退>` 兼容老客户端，
    `filename*=UTF-8''<百分号编码>` 才是真正展示给用户的名字 —— HTTP header 只能放
    latin-1，中文直接塞进去在部分客户端会乱码甚至报错。
    """

    def test_safe_filename_part_cleans_input(self):
        from app.utils.helpers import safe_filename_part

        assert safe_filename_part("SMC/华东") == "SMC 华东"        # 路径符号不落地
        assert safe_filename_part('a:b*c?d"e<f>g|h\\i') == "a b c d e f g h i"
        assert safe_filename_part("  ..  ") == ""
        assert safe_filename_part(None, "fallback") == "fallback"
        assert len(safe_filename_part("长" * 100)) == 40            # 超长截断

    def test_quotation_export_filename_is_number_and_title(self, db):
        """报价单文件名 = 编号 + 标题（不含「报价单」前缀与客户名）。

        用真实数据形态断言：客户叫「SMC」、标题就叫「SMC-会议室环境检测」。
        编号本身以 QT 开头（quotation），类型语义已在里面，不该再写「报价单_」；
        客户信息也已经写在标题里，不该再拼一遍客户名（R43/R44）。
        """
        from urllib.parse import quote

        qt = client.post("/product-db/api/quotations",
                         json={"title": "SMC-会议室环境检测", "client_name": "SMC"}).json()["quotation"]
        cd = client.get(f"/product-db/api/quotations/{qt['id']}/export-xlsx").headers["content-disposition"]
        assert cd.startswith("attachment;")
        assert f'filename="quotation_{qt["id"]}.xlsx"' in cd                      # ASCII 回退
        assert quote(f"{qt['quote_number']}_SMC-会议室环境检测", safe="") in cd
        assert quote("报价单", safe="") not in cd                                 # QT 已含该语义
        assert cd.count("SMC") == 1                                               # 客户名不得再拼一遍
        assert cd.index(quote(qt["quote_number"], safe="")) < cd.index("SMC")     # 标识在前

    def test_bom_export_filename_is_id_and_name(self, db):
        """BOM 文件名 = id + 方案名（用真实数据形态：方案名已含「客户-项目」）"""
        from urllib.parse import quote

        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)
        sol = Solution(name="麦当劳广州-空调集控", client_name="麦当劳（广州）")
        db.add(sol)
        db.commit()
        db.refresh(sol)
        db.add(SolutionItem(solution_id=sol.id, product_id=p.id, quantity=1,
                            unit_price=100, discount_rate=100))
        db.commit()

        res = client.get(f"/product-db/api/solutions/{sol.id}/bom-snapshot/export-xlsx")
        assert res.status_code == 200
        cd = res.headers["content-disposition"]
        assert quote(f"BOM_id{sol.id}_麦当劳广州-空调集控", safe="") in cd
        assert quote("麦当劳（广州）", safe="") not in cd        # 客户名不得再拼一遍
        assert f'filename="bom_solution_{sol.id}.xlsx"' in cd

    def test_product_export_filename_has_date(self, db):
        """产品清单是全库导出、没有客户/项目维度，用日期区分"""
        from datetime import datetime
        from urllib.parse import quote

        cd = client.get("/product-db/api/products/export").headers["content-disposition"]
        stamp = datetime.now().strftime("%Y%m%d")
        assert f'filename="products_{stamp}.xlsx"' in cd
        assert quote(f"产品清单_{stamp}", safe="") in cd

    def test_spec_sheet_pdf_filename(self, db, monkeypatch):
        """PDF 分支的文件名要带产品名与型号；顺带证明中文不会让 header 编码报错"""
        import subprocess
        from urllib.parse import quote

        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, name="雷达传感器/室外", model="VS373-470M")

        def fake_run(cmd, **kwargs):
            with open(cmd[2], "wb") as f:      # argv: [wp, html, pdf, -e, utf-8]
                f.write(b"%PDF-1.4" + b"0" * 200)
            return type("R", (), {"returncode": 0})()

        # 直接给 weasyprint 路径，绕开 shutil.which（该 import 在函数内部，无法按模块属性 patch）
        monkeypatch.setattr("app.config.settings.WEASYPRINT_PATH", "/usr/bin/weasyprint")
        monkeypatch.setattr(subprocess, "run", fake_run)

        res = client.get(f"/product-db/api/products/{p.id}/spec-sheet")
        assert res.status_code == 200
        assert res.headers["content-type"] == "application/pdf"
        cd = res.headers["content-disposition"]
        assert quote("产品规格书_VS373-470M_雷达传感器 室外", safe="") in cd
        assert f'filename="spec-sheet-{p.id}.pdf"' in cd

        # 缺型号时退回 id{产品ID}，保证「标识在最前」这条口径始终成立
        p2 = _seed_product(db, category_id=cat.id, name="无型号产品", model=None)
        res2 = client.get(f"/product-db/api/products/{p2.id}/spec-sheet")
        assert res2.status_code == 200
        assert quote(f"产品规格书_id{p2.id}_无型号产品", safe="") in res2.headers["content-disposition"]

    def test_spec_sheet_model_skipped_when_already_in_name(self, db, monkeypatch):
        """规格书：型号若已整段出现在产品名里，就不再单独拼一次。

        产品侧这是常见形态（名称「WTS506 气象站」+ 型号「WTS506」）。规格书的标识
        允许「藏在名称里」—— 名称本身已经带了这个标识，再拼就是重复。
        """
        import subprocess
        from urllib.parse import quote

        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, name="WTS506 气象站", model="WTS506")

        def fake_run(cmd, **kwargs):
            with open(cmd[2], "wb") as f:
                f.write(b"%PDF-1.4" + b"0" * 200)
            return type("R", (), {"returncode": 0})()

        monkeypatch.setattr("app.config.settings.WEASYPRINT_PATH", "/usr/bin/weasyprint")
        monkeypatch.setattr(subprocess, "run", fake_run)

        res = client.get(f"/product-db/api/products/{p.id}/spec-sheet")
        assert res.status_code == 200
        cd = res.headers["content-disposition"]
        assert quote("产品规格书_WTS506 气象站", safe="") in cd
        assert cd.count("WTS506") == 1        # 型号不得再拼一遍


# ============================================================
# R41: 成本出口补齐 + 批量删除不留孤儿 + token 可撤销
# ============================================================
class TestR41CostLeaks:
    """R41 修掉的两个成本泄漏出口。

    背景：R37 建立的约定是「成本的所有出口都必须走 cost_visible」，但
      ① 报价单条目的**写接口**（POST/PUT items）原样回传 `product_snapshot`，
         而同文件的列表接口已裁剪 —— 列表裁了、写接口没裁，等于没裁；
      ② 报价单**导出**只读全局开关，不看按用户的三态覆盖，
         于是「单独放行成本的用户」导出时仍拿不到成本（反过来则泄漏）。
    """

    def _make_user(self, db, username="normal", can_view_cost=None):
        u = User(username=username, password_hash=hash_password("test123"), role="user",
                 can_view_cost=can_view_cost)
        db.add(u)
        db.commit()
        db.refresh(u)
        return u

    def _auth_for(self, user):
        return {"Authorization": f"Bearer {create_token(user.id, user.username, user.token_version)}"}

    def _set_field(self, db, field_name, user_visible):
        from app.models.field_setting import FieldSetting
        fs = db.query(FieldSetting).filter_by(field_name=field_name).first()
        if fs:
            fs.user_visible = user_visible
        else:
            db.add(FieldSetting(field_name=field_name, user_visible=user_visible))
        db.commit()
        from app.services.field_visibility import _cache
        _cache["ts"] = 0
        _cache["data"] = None

    def _seed_quotation_with_item(self, db, user, product):
        qid = client.post("/product-db/api/quotations", json={"title": "T"},
                          headers=self._auth_for(user)).json()["quotation"]["id"]
        res = client.post(f"/product-db/api/quotations/{qid}/items",
                          json={"product_id": product.id, "quantity": 1, "unit_price": 50},
                          headers=self._auth_for(user))
        assert res.status_code == 201
        return qid, res

    def test_add_item_response_hides_cost_price(self, db):
        """新增条目的**响应体**不能带成本（此前只有随后的 GET 会裁剪）"""
        self._set_field(db, "cost_price", False)
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, base_price=100, cost_price=55.5)
        user = self._make_user(db)

        _, res = self._seed_quotation_with_item(db, user, p)
        assert "cost_price" not in res.json()["item"]["product_snapshot"]

        # 管理员仍要能看到（不能过度裁剪）
        admin = db.query(User).filter_by(username="admin").first()
        _, res = self._seed_quotation_with_item(db, admin, p)
        assert res.json()["item"]["product_snapshot"]["cost_price"] == 55.5

    def test_update_item_response_hides_cost_price(self, db):
        self._set_field(db, "cost_price", False)
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, base_price=100, cost_price=55.5)
        user = self._make_user(db)
        qid, res = self._seed_quotation_with_item(db, user, p)
        item_id = res.json()["item"]["id"]

        res = client.put(f"/product-db/api/quotations/{qid}/items/{item_id}",
                         json={"quantity": 3}, headers=self._auth_for(user))
        assert res.status_code == 200
        assert "cost_price" not in res.json()["item"]["product_snapshot"]

    def test_export_follows_per_user_cost_override(self, db):
        """导出必须走 cost_visible：按用户放行时要有成本列，按用户禁止时不能有"""
        import openpyxl
        from io import BytesIO

        self._set_field(db, "cost_price", False)          # 全局：隐藏
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, base_price=100, cost_price=66.6)
        allowed = self._make_user(db, "cost-allowed", can_view_cost=True)
        denied = self._make_user(db, "cost-denied", can_view_cost=False)

        def _cost_cell(user):
            qid, _ = self._seed_quotation_with_item(db, user, p)
            exp = client.get(f"/product-db/api/quotations/{qid}/export-xlsx",
                             headers=self._auth_for(user))
            assert exp.status_code == 200
            ws = openpyxl.load_workbook(BytesIO(exp.content)).active
            return ws.cell(row=4, column=13).value        # M 列 = 成本

        # 全局隐藏 + 按用户放行 → 必须看到（R37 的覆盖在这条路径上曾经失效）
        assert _cost_cell(allowed) == pytest.approx(66.6)

        # 全局打开 + 按用户禁止 → 必须看不到
        self._set_field(db, "cost_price", True)
        assert _cost_cell(denied) in (None, "")

    def test_export_omits_cost_header_entirely_when_hidden(self, db):
        """看不到成本时，导出的 xlsx 里不该出现「成本」两个字（R46）。

        只留空值、却保留「成本」表头，会让表格看起来像「有这一列但数据漏了」。
        报价单与产品清单统一为整列不输出（BOM 两个分支在 R45 已如此）。
        """
        import openpyxl
        from io import BytesIO

        self._set_field(db, "cost_price", False)
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, base_price=100, cost_price=66.6)
        denied = self._make_user(db, "no-cost-user", can_view_cost=False)
        admin = db.query(User).filter_by(username="admin").first()

        def _all_cells(ws):
            return [ws.cell(row=r, column=c).value
                    for r in range(1, ws.max_row + 1) for c in range(1, ws.max_column + 1)]

        # 报价单导出：表头、数据、列宽都不该有成本列
        qid, _ = self._seed_quotation_with_item(db, denied, p)
        ws = openpyxl.load_workbook(BytesIO(client.get(
            f"/product-db/api/quotations/{qid}/export-xlsx",
            headers=self._auth_for(denied)).content)).active
        assert ws.cell(row=3, column=13).value is None
        assert ws.cell(row=4, column=13).value is None
        assert "成本" not in _all_cells(ws)
        # 空列也不留成本列专用宽度（openpyxl 对未设置的列返回默认宽度，所以与常量比）
        from app.utils.excel_style import COLUMN_WIDTHS
        assert ws.column_dimensions["M"].width != COLUMN_WIDTHS["M"]

        # 产品清单导出
        ws = openpyxl.load_workbook(BytesIO(client.get(
            "/product-db/api/products/export", headers=self._auth_for(denied)).content)).active
        assert ws.cell(row=3, column=13).value is None
        assert "成本" not in _all_cells(ws)

        # 管理员仍应看到「成本」表头（不能过度裁剪）
        ws = openpyxl.load_workbook(BytesIO(client.get(
            f"/product-db/api/quotations/{qid}/export-xlsx",
            headers=self._auth_for(admin)).content)).active
        assert ws.cell(row=3, column=13).value == "成本"


class TestR41BatchDeleteNoOrphans:
    """批量删除必须级联清掉子行。

    原来用 `db.query(...).delete()`（bulk DELETE），不触发 ORM cascade，而 SQLite 的
    外键约束在生产未启用 → 子行永留库中。危害不止占空间：SQLite 的 INTEGER PRIMARY KEY
    会复用被删行的 rowid，新单据拿到同一 id 时会把历史孤儿子行「认领」进新单据。
    """

    @staticmethod
    def _orphan_count(db, table, column, parent):
        from sqlalchemy import text
        return db.execute(text(
            f"SELECT COUNT(*) FROM {table} WHERE {column} IS NOT NULL "
            f"AND {column} NOT IN (SELECT id FROM {parent})"
        )).scalar()

    def test_solution_batch_delete_removes_children(self, db, auth_headers):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)
        sol = Solution(name="待删方案")
        db.add(sol)
        db.commit()
        db.refresh(sol)
        db.add(SolutionItem(solution_id=sol.id, product_id=p.id, quantity=1, unit_price=10))
        db.add(SolutionBOMSnapshot(solution_id=sol.id, snapshot={"A1": {"v": "x"}}))
        db.commit()

        res = client.post("/product-db/api/solutions/batch-delete",
                          json={"ids": [sol.id]}, headers=auth_headers)
        assert res.status_code == 200
        assert res.json()["deleted"] == 1

        db.expire_all()
        assert self._orphan_count(db, "solution_items", "solution_id", "solutions") == 0
        assert self._orphan_count(db, "solution_bom_snapshots", "solution_id", "solutions") == 0

    def test_quotation_batch_delete_removes_children(self, db, auth_headers):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, base_price=10)
        qt = client.post("/product-db/api/quotations", json={"title": "待删报价单"},
                         headers=auth_headers).json()["quotation"]
        db.add(QuotationItem(quotation_id=qt["id"], product_id=p.id,
                             product_snapshot={"name": p.name}, quantity=1, unit_price=10))
        db.commit()

        res = client.post("/product-db/api/quotations/batch-delete",
                          json={"ids": [qt["id"]]}, headers=auth_headers)
        assert res.status_code == 200
        assert res.json()["deleted"] == 1

        db.expire_all()
        assert self._orphan_count(db, "quotation_items", "quotation_id", "quotations") == 0


class TestR41TokenRevocation:
    """JWT 撤销：登出 / 改密后，旧 token 必须立即失效。

    在此之前全仓没有登出接口、也没有任何撤销手段 —— 密码被改或被重置后，
    旧 token 仍能用满 JWT_EXPIRE_MINUTES（默认 24h）。
    """

    def _make_user(self, db, username="tokuser"):
        u = User(username=username, password_hash=hash_password("test123"), role="user")
        db.add(u)
        db.commit()
        db.refresh(u)
        return u

    def _auth_for(self, user):
        return {"Authorization": f"Bearer {create_token(user.id, user.username, user.token_version)}"}

    def test_logout_invalidates_existing_token(self, db):
        u = self._make_user(db)
        headers = self._auth_for(u)
        assert client.get("/product-db/api/auth/me", headers=headers).status_code == 200

        assert client.post("/product-db/api/auth/logout", headers=headers).status_code == 200
        assert client.get("/product-db/api/auth/me", headers=headers).status_code == 401

        # 重新登录后拿到的新 token 必须可用
        res = client.post("/product-db/api/auth/login",
                          json={"username": u.username, "password": "test123"})
        assert res.status_code == 200
        new_headers = {"Authorization": f"Bearer {res.json()['token']}"}
        assert client.get("/product-db/api/auth/me", headers=new_headers).status_code == 200

    def test_password_change_invalidates_existing_token(self, db):
        u = self._make_user(db)
        headers = self._auth_for(u)

        res = client.put("/product-db/api/auth/profile",
                         json={"current_password": "test123", "password": "newpass123"},
                         headers=headers)
        assert res.status_code == 200
        assert client.get("/product-db/api/auth/me", headers=headers).status_code == 401

    def test_admin_reset_invalidates_existing_token(self, db, auth_headers):
        u = self._make_user(db)
        headers = self._auth_for(u)
        assert client.get("/product-db/api/auth/me", headers=headers).status_code == 200

        res = client.put(f"/product-db/api/admin/users/{u.id}/password",
                         json={"password": "resetpass123"}, headers=auth_headers)
        assert res.status_code == 200
        assert client.get("/product-db/api/auth/me", headers=headers).status_code == 401

    def test_legacy_token_without_ver_still_valid(self, db):
        """升级不能把已登录的人踢下线：老 token 没有 ver 字段，按 0 比对"""
        import jwt
        from datetime import datetime, timedelta, timezone
        from app.config import settings

        u = self._make_user(db)
        payload = {"sub": str(u.id), "username": u.username,
                   "exp": datetime.now(timezone.utc) + timedelta(minutes=5)}
        legacy = jwt.encode(payload, settings.SECRET_KEY, algorithm=settings.JWT_ALGORITHM)

        res = client.get("/product-db/api/auth/me", headers={"Authorization": f"Bearer {legacy}"})
        assert res.status_code == 200
        assert res.json()["user"]["username"] == u.username


class TestBOMColumnLayout:
    """BOM 导出的两种来源（编辑器快照 / 兜底生成）必须用同一套列布局，成本固定在 J（R45）。

    历史上兜底分支硬套了报价单的 12 列（成本在 M），于是同一份 BOM「有没有在编辑器里
    保存过」会导出两种列结构；而存储、`_keep_server_cost`、`_COST_COLUMN` 一律按 J 列
    认定成本 —— 布局与判定口径不一致。
    """

    def _seed_solution(self, db):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, name="雷达传感器", model="VS373",
                          base_price=100, cost_price=66.6)
        sol = Solution(name="某园区方案")
        db.add(sol)
        db.commit()
        db.refresh(sol)
        db.add(SolutionItem(solution_id=sol.id, product_id=p.id, quantity=2,
                            unit_price=100, discount_rate=100))
        db.commit()
        return p, sol

    def test_basic_export_uses_snapshot_column_layout(self, db):
        """没有快照（走兜底生成）时，列头与成本列位置仍与编辑器快照一致"""
        import openpyxl
        from io import BytesIO

        _, sol = self._seed_solution(db)
        res = client.get(f"/product-db/api/solutions/{sol.id}/bom-snapshot/export-xlsx")
        assert res.status_code == 200
        ws = openpyxl.load_workbook(BytesIO(res.content)).active

        headers = [ws.cell(row=3, column=c).value for c in range(1, 11)]
        assert headers == ["#", "产品名称", "型号/SKU", "功能描述", "数量",
                           "单价", "折扣%", "小计", "备注", "成本"]
        assert ws.cell(row=4, column=1).value == 1              # A 序号
        assert ws.cell(row=4, column=5).value == 2              # E 数量
        assert ws.cell(row=4, column=6).value == 100            # F 单价
        assert ws.cell(row=4, column=8).value == "=E4*F4*G4/100"   # H 小计 = 数量 × 单价 × 折扣%
        assert ws.cell(row=4, column=10).value == 66.6          # J 成本
        # 第 11 列不该再有任何内容 —— 旧的 12 列布局会在这里留「图片」列
        assert ws.cell(row=3, column=11).value is None

    def test_basic_export_hides_cost_column_entirely_for_user(self, db):
        """看不到成本时整列不输出（连表头也不写），与快照分支的剥离口径一致"""
        import openpyxl
        from io import BytesIO
        from app.models.field_setting import FieldSetting
        from app.services.field_visibility import _cache

        db.add(FieldSetting(field_name="cost_price", user_visible=False))
        db.commit()
        _cache["ts"] = 0
        _cache["data"] = None

        _, sol = self._seed_solution(db)
        u = User(username="bomuser", password_hash=hash_password("test123"), role="user")
        db.add(u)
        db.commit()
        db.refresh(u)
        headers = {"Authorization": f"Bearer {create_token(u.id, u.username, u.token_version)}"}

        res = client.get(f"/product-db/api/solutions/{sol.id}/bom-snapshot/export-xlsx",
                         headers=headers)
        assert res.status_code == 200
        ws = openpyxl.load_workbook(BytesIO(res.content)).active
        assert ws.cell(row=3, column=9).value == "备注"          # 第 9 列仍在
        assert ws.cell(row=3, column=10).value is None           # 成本表头不写
        assert ws.cell(row=4, column=10).value is None           # 成本值也不写

        # 复位缓存，避免影响后续用例（FieldSetting 行会随每个用例的 drop_all 清掉）
        _cache["ts"] = 0
        _cache["data"] = None

    def test_basic_export_subtotal_includes_discount(self, db):
        """兜底分支的「小计」必须含折扣，与快照分支 `qty*price*discount/100` 同口径（R48）。

        原先是 `=E*F`（不含折扣），而合计行的大写金额来自含折扣的 `sol.total_price`
        → 同一张表里 `SUM(H)` 与大写金额会自相矛盾。
        """
        import openpyxl
        from io import BytesIO

        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, base_price=100)
        sol = Solution(name="带折扣方案", total_price=900)      # 10 × 100 × 90% = 900
        db.add(sol)
        db.commit()
        db.refresh(sol)
        db.add(SolutionItem(solution_id=sol.id, product_id=p.id, quantity=10,
                            unit_price=100, discount_rate=90))
        db.commit()

        res = client.get(f"/product-db/api/solutions/{sol.id}/bom-snapshot/export-xlsx")
        assert res.status_code == 200
        ws = openpyxl.load_workbook(BytesIO(res.content)).active
        assert ws.cell(row=4, column=7).value == 90                    # G 折扣%
        assert ws.cell(row=4, column=8).value == "=E4*F4*G4/100"       # H 小计含折扣
        assert "=SUM(H4:H4)" in str(ws.cell(row=5, column=8).value)    # 合计落在小计列（1 条数据 → 第 5 行）


class TestR48Hardening:
    """R48 修的三条高优先项里，与鉴权/提示词有关的两条。"""

    def test_agent_prompt_uses_header_auth_not_query_token(self):
        """agent 提示词必须用请求头传 token（R48）。

        后端只允许 **GET** 用 query-string token（`auth.py` 的 method 检查），
        而写操作（POST/PUT/DELETE）带 `?token=` 一律 401 —— 原提示词恰好要求 agent 这么做，
        等于让 AI 的建/改/删方案与建报价单必然失败。
        """
        import re
        from app.routers.admin_routes import _PROMPT_DEFAULTS

        prompt = _PROMPT_DEFAULTS["agent_prompt"]
        assert "Authorization: Bearer {{TOKEN}}" in prompt
        assert not re.search(r"[/?&]token=\{\{TOKEN\}\}", prompt), "URL 里不该再带 token 占位符"

    def test_admin_cannot_change_or_reset_own_password(self, db, auth_headers):
        """管理员不能通过管理接口改/重置**自己**的密码（R48）。

        那会把自己刚签发的 token 作废（token_version +1），而管理页用的 adminApi 是原生
        fetch、不接 api() 的 401 跳转 → 自己把自己锁在报错页。改自己密码要走「个人信息」。
        """
        admin = db.query(User).filter_by(username="admin").first()

        res = client.put(f"/product-db/api/admin/users/{admin.id}/password",
                         json={"password": "whatever123"}, headers=auth_headers)
        assert res.status_code == 400
        assert "个人信息" in res.json()["detail"]

        res = client.put(f"/product-db/api/admin/users/{admin.id}",
                         json={"password": "whatever123"}, headers=auth_headers)
        assert res.status_code == 400

    def test_admin_can_still_reset_other_users_password(self, db, auth_headers):
        """改别人的密码不受影响，且确实作废了对方的 token（避免为修自锁而废掉正常功能）"""
        victim = User(username="victim", password_hash=hash_password("test123"), role="user")
        db.add(victim)
        db.commit()
        db.refresh(victim)
        victim_token = create_token(victim.id, victim.username, victim.token_version)

        res = client.put(f"/product-db/api/admin/users/{victim.id}/password",
                         json={"password": "newpass1234"}, headers=auth_headers)
        assert res.status_code == 200
        assert client.get("/product-db/api/auth/me",
                          headers={"Authorization": f"Bearer {victim_token}"}).status_code == 401


class TestDiscountZero:
    """折扣率 0 是合法值（免费/赠品），不能被 `or 100` 当成「未设置」按 100% 算（R49）。

    历史缺陷：明细金额、方案/报价单合计、导出表格、AI 建单都写成 `discount_rate or 100`，
    于是 0 被当成未填 → 按 100% 计价。判定现统一走 `helpers.discount_percent`。
    """

    def test_discount_percent_helper_basics(self):
        from app.utils.helpers import discount_percent

        assert discount_percent(0) == 0              # 0 必须保留 —— 核心回归点
        assert discount_percent("0") == 0            # 导入表格里可能是字符串
        assert discount_percent(None) == 100         # 未填才取默认
        assert discount_percent("") == 100
        assert discount_percent("abc") == 100        # 脏数据不抛异常
        assert discount_percent(90) == 90
        assert discount_percent(0, default=1) == 0   # default 只对空值生效

    def test_solution_item_with_zero_discount(self, db, auth_headers):
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, base_price=1000)
        sol = Solution(name="赠品方案")
        db.add(sol)
        db.commit()
        db.refresh(sol)

        res = client.post(f"/product-db/api/solutions/{sol.id}/items",
                          json={"product_id": p.id, "quantity": 2, "unit_price": 1000,
                                "discount_rate": 0}, headers=auth_headers)
        assert res.status_code == 201

        db.expire_all()
        item = db.query(SolutionItem).filter_by(solution_id=sol.id).first()
        assert float(item.discount_rate) == 0, "折扣率 0 不能落库成 100"
        # SolutionItem 没有 amount 列，金额只体现在方案合计上
        assert float(db.get(Solution, sol.id).total_price or 0) == 0

    def test_quotation_item_with_zero_discount_and_export(self, db, auth_headers):
        import openpyxl
        from io import BytesIO

        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, base_price=1000)
        qt = Quotation(title="赠品报价", client_name="客户A")
        db.add(qt)
        db.commit()
        db.refresh(qt)

        res = client.post(f"/product-db/api/quotations/{qt.id}/items",
                          json={"product_id": p.id, "quantity": 2, "unit_price": 1000,
                                "discount_rate": 0}, headers=auth_headers)
        assert res.status_code == 201

        db.expire_all()
        item = db.query(QuotationItem).filter_by(quotation_id=qt.id).first()
        assert float(item.discount_rate) == 0
        assert float(item.amount or 0) == 0
        assert float(db.get(Quotation, qt.id).total_amount or 0) == 0

        # 导出表格里折扣列（I）应为 0（旧代码会写 100 → 与库里 0 不符）
        res = client.get(f"/product-db/api/quotations/{qt.id}/export-xlsx", headers=auth_headers)
        assert res.status_code == 200
        ws = openpyxl.load_workbook(BytesIO(res.content)).active
        assert ws.cell(row=4, column=9).value == 0


class TestQuantityZero:
    """数量 0 是合法值（本次不采购但保留该行），不能被 `or 1` 悄悄改成 1（R50）。

    与折扣 0 同族：只有 None/空/脏数据才回退默认值，0 原样保留。
    """

    def test_number_or_helper_basics(self):
        from app.utils.helpers import discount_percent, number_or

        assert number_or(0, 1) == 0            # 0 必须保留 —— 核心回归点
        assert number_or("0", 1) == 0          # BOM 表格读到的可能是字符串
        assert number_or(None, 1) == 1         # 空才取默认
        assert number_or("", 1) == 1
        assert number_or("abc", 1) == 1        # 脏数据不抛异常
        assert number_or(5, 1) == 5
        # discount_percent 复用同一实现，语义不漂移
        assert discount_percent(0) == 0
        assert discount_percent(None) == 100

    def test_bom_save_keeps_zero_quantity_row(self, db, auth_headers):
        """BOM 编辑保存数量 0：该行保留、金额 0，且不影响其他行合计"""
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id, base_price=1000)
        qt = Quotation(title="BOM 数量0", client_name="客户B")
        db.add(qt)
        db.commit()
        db.refresh(qt)

        res = client.put(f"/product-db/api/quotations/{qt.id}/bom", json={"rows": [
            {"name": p.name, "sku": "SKU-A", "model": p.model, "qty": 0,
             "price": 1000, "discount": 100, "description": "", "remark": "本次不采购"},
            {"name": p.name, "sku": "SKU-B", "model": p.model, "qty": 2,
             "price": 1000, "discount": 100, "description": "", "remark": ""},
        ]}, headers=auth_headers)
        assert res.status_code == 200
        assert res.json()["total"] == 2000

        db.expire_all()
        items = (db.query(QuotationItem).filter_by(quotation_id=qt.id)
                 .order_by(QuotationItem.sort_order).all())
        assert len(items) == 2, "数量 0 的行必须保留，不能被过滤掉"
        assert float(items[0].quantity) == 0, "数量 0 不能落库成 1"
        assert float(items[0].amount or 0) == 0
        assert float(items[1].quantity) == 2


class TestConfigConsistency:
    """库路径只有一个来源（R52）。

    R51 事故根因：默认 DATABASE_URL 用 os.path.expanduser('~') 拼家目录，
    systemd 启用 ProtectHome 隐藏 /home 后，服务能启动但所有查库接口 500。
    """

    def test_default_database_url_is_inside_backend_dir(self, monkeypatch):
        from app import config as cfg

        # 必须先清掉环境变量：pydantic-settings 里 env 的优先级高于代码默认值，
        # 而本文件顶部已把 DATABASE_URL 指向了临时测试库，否则拿到的是那个值
        monkeypatch.delenv("DATABASE_URL", raising=False)
        monkeypatch.delenv("DATABASE_PATH", raising=False)

        # _env_file=None：也不读 .env，这样拿到的才是代码里的默认值
        defaults = cfg.Settings(_env_file=None, SECRET_KEY="x" * 32)
        expected = os.path.join(cfg._BACKEND_DIR, "product_db.db")
        assert defaults.DATABASE_URL == f"sqlite:///{expected}"
        assert "/home/" not in defaults.DATABASE_URL
        assert "/root/" not in defaults.DATABASE_URL

    def test_db_filesystem_path_matches_database_url(self):
        """reporting 用的 DB_FILESYSTEM_PATH 必须与 engine 用的 DATABASE_URL 指向同一文件。

        以前 DB_FILESYSTEM_PATH 优先取 DATABASE_PATH，会出现「engine 读写 A 路径、
        agent 却报告 B 路径」的不一致。
        """
        from app.config import DB_FILESYSTEM_PATH, settings

        assert DB_FILESYSTEM_PATH == settings.DATABASE_URL.replace("sqlite:///", "", 1)


class TestR55Security:
    """R55：agent 审批接口越权（S1/S2）+ 用户可控 URL 的协议白名单（S3）。"""

    def _make_user(self, db, username, role="user"):
        u = User(username=username, password_hash=hash_password("test123"), role=role)
        db.add(u)
        db.commit()
        db.refresh(u)
        return u

    @staticmethod
    def _auth(u):
        return {"Authorization": f"Bearer {create_token(u.id, u.username, u.token_version)}"}

    def test_approvals_are_scoped_to_owner(self, db, auth_headers):
        """S1：普通用户只能看到自己的待审批任务，管理员看全部"""
        from app.services.approval_manager import approval_manager

        owner = self._make_user(db, "r55_owner")
        other = self._make_user(db, "r55_other")
        task = approval_manager.create(
            tool_name="create_quotation", tool_label="创建报价单",
            tool_input={"customer": "owner 的客户"}, summary="s", details={}, user_id=owner.id)
        try:
            tasks = client.get("/product-db/api/agent/approvals", headers=auth_headers).json()["tasks"]
            assert any(t["task_id"] == task.task_id for t in tasks), "管理员应能看到全部"

            tasks = client.get("/product-db/api/agent/approvals", headers=self._auth(owner)).json()["tasks"]
            assert any(t["task_id"] == task.task_id for t in tasks), "所有者应能看到自己的"

            tasks = client.get("/product-db/api/agent/approvals", headers=self._auth(other)).json()["tasks"]
            assert all(t["task_id"] != task.task_id for t in tasks), "他人任务（含 tool_input 明细）不得可见"
        finally:
            approval_manager._tasks.pop(task.task_id, None)

    def test_test_approval_requires_admin(self, db, auth_headers):
        """S2：test-approval 只允许管理员（会往审批队列里塞假任务）"""
        from app.services.approval_manager import approval_manager

        normal = self._make_user(db, "r55_normal")
        assert client.post("/product-db/api/agent/test-approval",
                           headers=self._auth(normal)).status_code == 403

        res = client.post("/product-db/api/agent/test-approval", headers=auth_headers)
        assert res.status_code == 200
        approval_manager._tasks.pop(res.json()["task_id"], None)

    def test_is_safe_user_url(self):
        """S3：用户填写的链接只允许 http/https（防 javascript:/data: 变成 XSS 跳板）"""
        from app.utils.security import is_safe_user_url

        assert is_safe_user_url("https://example.com/a?b=1")
        assert is_safe_user_url("http://192.168.1.10:8080/portal")   # 内网链接是正常需求，不拦
        assert not is_safe_user_url("javascript:alert(1)")
        assert not is_safe_user_url("data:text/html,<script>alert(1)</script>")
        assert not is_safe_user_url("ftp://example.com/x")
        assert not is_safe_user_url("example.com")                    # 无协议
        assert not is_safe_user_url("https://")                       # 无主机
        assert not is_safe_user_url("")

    def test_product_url_schema_rejects_executable_scheme(self):
        """S3：product_url 在 schema 层就拦掉可执行协议（前端会直接 window.open 它）"""
        from app.schemas.product import ProductCreate

        ok = ProductCreate(name="x", category_id=1, product_url="https://example.com/p")
        assert ok.product_url == "https://example.com/p"
        with pytest.raises(Exception):
            ProductCreate(name="x", category_id=1, product_url="javascript:alert(1)")

    def test_file_link_rejects_executable_scheme(self, db, auth_headers):
        """S3：附件外链同样只允许 http/https"""
        cat = _seed_category(db)
        p = _seed_product(db, category_id=cat.id)

        res = client.post(f"/product-db/api/products/{p.id}/links",
                          json={"link_url": "javascript:alert(1)", "label": "x"}, headers=auth_headers)
        assert res.status_code == 400
        assert "http" in res.json()["detail"]

        res = client.post(f"/product-db/api/products/{p.id}/links",
                          json={"link_url": "https://example.com/doc", "label": "手册"}, headers=auth_headers)
        assert res.status_code in (200, 201)


class TestR56Consistency:
    """R56：D1 品类防成环 · D2 报价单取号 · D3 GET 快照不写库 · D4 导入价格容错。"""

    def test_category_cycle_guard(self, db, auth_headers):
        """D1：不能把品类挂到自身或自己的后代下（成环后取后代/建树会递归崩溃）"""
        from app.services.product_category_helper import (
            get_category_descendants, would_create_category_cycle)

        root = _seed_category(db, name="R56 根", slug="r56-root")
        mid = Category(name="R56 中", slug="r56-mid", parent_id=root.id, level=1)
        db.add(mid)
        db.commit()
        db.refresh(mid)
        leaf = Category(name="R56 叶", slug="r56-leaf", parent_id=mid.id, level=2)
        db.add(leaf)
        db.commit()
        db.refresh(leaf)

        assert set(get_category_descendants(db, root.id)) == {root.id, mid.id, leaf.id}
        assert would_create_category_cycle(db, root.id, root.id) is True
        assert would_create_category_cycle(db, root.id, leaf.id) is True    # 挂到后代 → 成环
        assert would_create_category_cycle(db, leaf.id, root.id) is False   # 挂到祖先 → 合法
        assert would_create_category_cycle(db, root.id, None) is False

        r = client.put(f"/product-db/api/categories/{root.id}",
                       json={"parent_id": root.id}, headers=auth_headers)
        assert r.status_code == 400
        r = client.put(f"/product-db/api/categories/{root.id}",
                       json={"parent_id": leaf.id}, headers=auth_headers)
        assert r.status_code == 400
        # 合法改动仍要能通过（别把功能一起拦掉）
        r = client.put(f"/product-db/api/categories/{leaf.id}",
                       json={"parent_id": root.id}, headers=auth_headers)
        assert r.status_code == 200

    def test_descendants_survive_existing_cycle(self, db):
        """D1：库里已存在环（历史脏数据）时，取后代不能无限递归"""
        from app.services.product_category_helper import get_category_descendants

        a = Category(name="R56 环A", slug="r56-a")
        b = Category(name="R56 环B", slug="r56-b")
        db.add_all([a, b])
        db.commit()
        db.refresh(a)
        db.refresh(b)
        a.parent_id = b.id
        b.parent_id = a.id
        db.commit()

        ids = get_category_descendants(db, a.id)    # 关键：不能 RecursionError
        assert a.id in ids and b.id in ids

    def test_quote_number_takes_max_sequence(self, db, auth_headers):
        """D2：取号取当天最大序号+1（旧实现按 id 倒序，遇乱序就会重号撞唯一约束）"""
        from datetime import datetime, timezone

        today = datetime.now(timezone.utc).strftime("%Y%m%d")
        for num in ("001", "003", "002"):
            db.add(Quotation(quote_number=f"QT-{today}-{num}", title=f"R56 乱序 {num}"))
        db.commit()

        res = client.post("/product-db/api/quotations", json={"title": "R56 新单"}, headers=auth_headers)
        assert res.status_code == 201
        assert res.json()["quotation"]["quote_number"] == f"QT-{today}-004"

    def test_get_bom_snapshot_does_not_duplicate(self, db, auth_headers):
        """D3：多次 GET 只落一条快照（唯一约束冲突会兜底复用，不再 500）"""
        from app.models.bom_template import SolutionBOMSnapshot

        sol = Solution(name="R56 快照幂等")
        db.add(sol)
        db.commit()
        db.refresh(sol)

        for _ in range(3):
            r = client.get(f"/product-db/api/solutions/{sol.id}/bom-snapshot", headers=auth_headers)
            assert r.status_code == 200
            assert r.json()["bom_snapshot"]["snapshot"] is not None

        db.expire_all()
        assert db.query(SolutionBOMSnapshot).filter_by(solution_id=sol.id).count() == 1

    def test_num_or_400_reports_row_number(self):
        """D4：非数字价格要报错并指出行号（原来整批 500 且不知哪一行）"""
        from fastapi import HTTPException as FastApiHTTPException
        from app.routers.product_import import _num_or_400

        assert _num_or_400(100, "价格", 1) == 100
        assert _num_or_400("1200.5", "价格", 1) == 1200.5
        # R75 改：不填存 NULL（未定价），与「填了 0 元」是两件事
        assert _num_or_400(None, "价格", 1) is None
        assert _num_or_400("", "成本", 3) is None
        assert _num_or_400(0, "价格", 1) == 0.0
        with pytest.raises(FastApiHTTPException) as exc:
            _num_or_400("面议", "价格", 2)
        assert "2 个数据行" in exc.value.detail
        assert "面议" in exc.value.detail


class TestR57ForeignKeys:
    """R57：外键强制开启（PRAGMA foreign_keys=ON）+ 迁移补的 ON DELETE 语义。

    以前 SQLite 默认不校验外键，模型里写的 ondelete 全是摆设：删父行只会静默留孤儿
    （生产实测 787 行违规）。开启强制的前提是存量违规已清、缺失的 ON DELETE 已补齐
    （迁移 e0f1a2b3c4d5），否则「删用户」会从留孤儿变成直接 500。
    """

    def test_foreign_keys_pragma_is_on(self, db):
        """守卫：database.py 的连接事件必须真的把 PRAGMA 打开了"""
        from sqlalchemy import text

        assert db.execute(text("PRAGMA foreign_keys")).scalar() == 1

    def test_product_with_unknown_category_is_rejected(self, db):
        """行为守卫：引用不存在的品类要立刻报错（不再静默写入孤儿）"""
        from sqlalchemy.exc import IntegrityError

        db.add(Product(name="R57 孤儿品类", category_id=999999))
        with pytest.raises(IntegrityError):
            db.commit()
        db.rollback()

    def test_deleting_user_nulls_ownership_and_keeps_records(self, db):
        """删用户：归属列 SET NULL 保留业务记录，而不是报错或留孤儿"""
        owner = User(username="r57-owner", password_hash=hash_password("x"), role="user")
        db.add(owner)
        db.commit()
        db.refresh(owner)

        sol = Solution(name="R57 方案", created_by=owner.id)
        db.add(sol)
        db.commit()
        db.refresh(sol)
        sol_id = sol.id

        db.delete(owner)
        db.commit()
        db.expire_all()

        kept = db.get(Solution, sol_id)
        assert kept is not None, "删用户不该连带删掉方案"
        assert kept.created_by is None

    def test_deleting_bom_template_cascades_snapshot(self, db):
        """删模板：快照随之级联删除（快照脱离模板没有意义）"""
        tmpl = BOMTemplate(name="R57 模板", snapshot="{}")
        db.add(tmpl)
        db.commit()
        db.refresh(tmpl)

        sol = Solution(name="R57 模板方案")
        db.add(sol)
        db.commit()
        db.refresh(sol)

        snap = SolutionBOMSnapshot(solution_id=sol.id, template_id=tmpl.id, snapshot="{}")
        db.add(snap)
        db.commit()
        snap_id = snap.id

        db.delete(tmpl)
        db.commit()
        db.expire_all()

        assert db.get(SolutionBOMSnapshot, snap_id) is None

    def test_deleting_solution_cascades_children(self, db):
        """删方案：条目与快照级联清掉，不留孤儿"""
        sol = Solution(name="R57 级联方案")
        db.add(sol)
        db.commit()
        db.refresh(sol)

        cat = _seed_category(db, name="R57 品类", slug="r57-cat")
        prod = Product(name="R57 产品", category_id=cat.id)
        db.add(prod)
        db.commit()
        db.refresh(prod)

        db.add(SolutionItem(solution_id=sol.id, product_id=prod.id, quantity=2))
        db.add(SolutionBOMSnapshot(solution_id=sol.id, template_id=None, snapshot="{}"))
        db.commit()
        sol_id = sol.id

        db.delete(sol)
        db.commit()
        db.expire_all()

        assert db.query(SolutionItem).filter_by(solution_id=sol_id).count() == 0
        assert db.query(SolutionBOMSnapshot).filter_by(solution_id=sol_id).count() == 0
        assert db.get(Product, prod.id) is not None, "产品不该被连带删除"

    def test_create_product_with_unknown_category_returns_400(self, auth_headers):
        """表单里选的品类已被删：给 400 说明，不再撞外键报 500"""
        r = client.post("/product-db/api/products",
                        json={"name": "R57 无效品类", "category_id": 999999},
                        headers=auth_headers)
        assert r.status_code == 400
        assert "品类不存在" in str(r.json()["detail"])

    def test_update_product_with_deleted_manufacturer_returns_400(self, db, auth_headers):
        """下拉框打开期间制造商被删掉：保存要提示刷新，不能 500"""
        from app.models.dictionary import Manufacturer

        cat = _seed_category(db, name="R57 厂商品类", slug="r57-mfg-cat")
        mfg = Manufacturer(name="R57 厂商")
        db.add(mfg)
        db.commit()
        db.refresh(mfg)
        mfg_id = mfg.id

        prod = Product(name="R57 厂商产品", category_id=cat.id, manufacturer_id=mfg_id)
        db.add(prod)
        db.commit()
        db.refresh(prod)
        prod_id = prod.id

        db.delete(mfg)   # 模拟管理员删掉该制造商（产品会被 SET NULL）
        db.commit()

        r = client.put(f"/product-db/api/products/{prod_id}",
                       json={"manufacturer_id": mfg_id}, headers=auth_headers)
        assert r.status_code == 400
        assert "已被删除" in r.json()["detail"]

    def test_product_used_by_solution_returns_409_not_500(self, db, auth_headers):
        """产品被方案引用时 RESTRICT 已生效：接口要给 409 说清楚，不能抛 500"""
        cat = _seed_category(db, name="R57 引用品类", slug="r57-ref-cat")
        prod = Product(name="R57 被引用产品", category_id=cat.id)
        db.add(prod)
        db.commit()
        db.refresh(prod)

        sol = Solution(name="R57 引用方案")
        db.add(sol)
        db.commit()
        db.refresh(sol)
        db.add(SolutionItem(solution_id=sol.id, product_id=prod.id, quantity=1))
        db.commit()

        r = client.delete(f"/product-db/api/products/{prod.id}", headers=auth_headers)
        assert r.status_code == 409
        assert "1 个方案条目" in r.json()["detail"]
        assert db.get(Product, prod.id) is not None, "被引用的产品不该被删掉"

        # 移除引用后要能正常删
        db.query(SolutionItem).filter_by(product_id=prod.id).delete()
        db.commit()
        r = client.delete(f"/product-db/api/products/{prod.id}", headers=auth_headers)
        assert r.status_code == 200

    def test_deleting_last_category_in_use_returns_400(self, db, auth_headers):
        """最后一个品类还有产品归属时：不能删（原来会把 category_id 改成不存在的 1）"""
        cat = _seed_category(db, name="R57 唯一品类", slug="r57-only-cat")
        cat_id = cat.id
        prod = Product(name="R57 归属产品", category_id=cat_id)
        db.add(prod)
        db.commit()

        r = client.delete(f"/product-db/api/categories/{cat_id}", headers=auth_headers)
        assert r.status_code == 400
        assert db.get(Category, cat_id) is not None

        # 没产品归属时（或新建了接手品类后）仍可删
        db.delete(prod)
        db.commit()
        r = client.delete(f"/product-db/api/categories/{cat_id}", headers=auth_headers)
        assert r.status_code == 200
        db.expire_all()
        assert db.get(Category, cat_id) is None

